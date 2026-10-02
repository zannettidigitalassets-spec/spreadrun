"""Build the UAD 3.6 rule table from the GSE-published appendices.

    python3 scripts/uad/build_rules.py --spec-dir <folder with the GSE workbooks>

Inputs (downloaded from the GSE UAD pages, never committed: the repo is public):
  Appendix A-1  URAR Delivery Specification   (sheet "UAD Delivery Spec 1.4", "UAD Enumerations")
  Appendix H-1  URAR Compliance Rules         (sheet "UAD Compliance Rules v1.5")

Output: pylib/spreadrun_api/validators/uad/rules.json, our own machine-readable table of
data point paths, formats, supported enumerations, cardinality and the compliance rules we can
evaluate deterministically. MISMO definitions and the MISMO XSD are not read and not copied.

H-1 rule logic is plain English. This script translates the shapes that have one reading
(required, conditionally required, comparisons, instance counts, uniqueness, date formats). Every
rule it cannot translate is written to "notImplemented" with the reason, so coverage is explicit.
"""
import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'pylib' / 'spreadrun_api' / 'validators' / 'uad' / 'rules.json'
COVERAGE = ROOT / 'src' / 'content' / 'uad-coverage.json'

VA = ('MESSAGE/DOCUMENT_SETS/DOCUMENT_SET/DOCUMENTS/DOCUMENT/DEAL_SETS/DEAL_SET/DEALS/DEAL/SERVICES/SERVICE/'
      'VALUATION/VALUATION_RESPONSE/VALUATION_ANALYSES/VALUATION_ANALYSIS')
PROPERTY = VA + '/PROPERTIES/PROPERTY'
USE_COLUMNS = {
    'Subject Property': 'SubjectProperty', 'Land Comparable': 'LandComparable',
    'Property Analyzed Not Used': 'PropertyAnalyzedNotUsed', 'Rental Comparable': 'RentalComparable',
    'Gross Rent Multiplier Comparable': 'GrossRentMultiplierComparable', 'Sales Comparable': 'SalesComparable',
}
AFFECTED = {
    'subject': 'SubjectProperty', 'sales comparable': 'SalesComparable', 'land comparable': 'LandComparable',
    'rental comparable': 'RentalComparable', 'gross rent multiplier comparable': 'GrossRentMultiplierComparable',
    'property analyzed not used': 'PropertyAnalyzedNotUsed',
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sheet_rows(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    header = [(h or '').split('\n')[0].strip() for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


def clean(v):
    if v is None:
        return ''
    return str(v).replace('_x000D_', '').strip()


# ---------------------------------------------------------------- A-1: paths, formats, enumerations

def card_max(c):
    c = clean(c)
    m = re.fullmatch(r'(\d+):(\d+|∞)', c)
    if not m:
        return None, None
    return int(m.group(1)), (None if m.group(2) == '∞' else int(m.group(2)))


def build_spec(a1_rows):
    containers = {}
    bounds = defaultdict(list)
    contexts = defaultdict(lambda: defaultdict(int))
    datapoints = defaultdict(list)
    for r in a1_rows:
        path = clean(r['MISMO xPath']).rstrip('/')
        if not path:
            continue
        name, attr = clean(r['MISMO Data Point Name']), clean(r['MISMO Attribute Name'])
        lo, hi = card_max(r['Cardinality'])
        use = sorted(v for k, v in USE_COLUMNS.items() if clean(r.get(k)).lower() == 'x')
        if not name and not attr:
            containers.setdefault(path, {'min': None, 'max': None})
            if lo is not None:  # the opening row of a container block carries its cardinality
                bounds[path].append((lo, hi, tuple(use)))
                contexts[path][tuple(use)] += 1
            continue
        if name:
            enum_raw = clean(r['Supported Data Point Enumerations'])
            items = [(name, enum_raw)]
        else:
            # One cell can describe several attributes: "@ValuationUseType\n@xlink:label" with
            # enumerations "@ValuationUseType - SalesComparable\n@xlink:label - PROPERTY_n".
            attrs = [a.strip().lstrip('@') for a in attr.split('\n') if a.strip()]
            enum_raw = clean(r['Supported Attribute Enumerations'])
            if len(attrs) > 1:
                per = {}
                for line in enum_raw.split('\n'):
                    m = re.match(r'\s*@?([\w:]+)\s*-\s*(.*)$', line)
                    if m:
                        per[m.group(1)] = m.group(2).strip()
                items = [('@' + a, per.get(a, '')) for a in attrs]
            else:
                items = [('@' + attrs[0], enum_raw)]
        for leaf, raw in items:
            datapoints[path + '/' + leaf].append({
                'id': clean(r['Unique ID']),
                'fmt': clean(r['UAD Data Point Format Type']) if name else '',
                'det': clean(r['UAD Data Point Format Details']) if name else '',
                'enum': [e.strip() for e in raw.split('|') if e.strip()] if raw else [],
                'use': use,
                'req': clean(r['Conditionality']),
            })
    # A container listed on several rows is usually one row per kind of instance (MARKET_INVENTORY for
    # active listings, pending sales, ...), so the allowed total is the sum. That over-allows when rows
    # repeat per valuation use type, which only makes the cardinality check more lenient, never stricter.
    def total(bs):
        return None if any(b[1] is None for b in bs) else sum(b[1] for b in bs)

    for path, bs in bounds.items():
        containers[path] = {'min': min(b[0] for b in bs), 'max': total(bs)}
        uses = {u for b in bs for u in b[2]}
        if uses:  # under a valuation PROPERTY: the bound for that use type alone
            containers[path]['maxByUse'] = {u: total([b for b in bs if not b[2] or u in b[2]]) for u in sorted(uses)}
    # How many distinct business contexts share this path for one use type (e.g. DATA_SOURCE for
    # appraisal sources and for price trend sources). Required checks only run where this is 1.
    for path, groups in contexts.items():
        containers[path]['contexts'] = max(groups.values())
    # every ancestor of a data point is a known container even if A-1 has no row for it
    for key in list(datapoints):
        parts = key.split('/')[:-1]
        for i in range(1, len(parts) + 1):
            containers.setdefault('/'.join(parts[:i]), {'min': None, 'max': None})
    return containers, dict(datapoints)


# ---------------------------------------------------------------- name resolution

class Untranslatable(Exception):
    pass


def lcp(a, b):
    a, b = a.split('/'), b.split('/')
    n = 0
    while n < min(len(a), len(b)) and a[n] == b[n]:
        n += 1
    return n


class Resolver:
    def __init__(self, containers, datapoints):
        self.containers = containers
        self.by_name = defaultdict(set)
        for key in datapoints:
            self.by_name[key.rsplit('/', 1)[1]].add(key)

    def leaf(self, name, context, target):
        """Full A-1 path(s) of data point / attribute `name` nearest to `context` (then `target`)."""
        if name.startswith('@'):
            cands = self.by_name.get(name, set())
            if name == '@ValuationUseType':
                cands = {PROPERTY + '/@ValuationUseType'}
        else:
            cands = self.by_name.get(name, set())
        if not cands:
            raise Untranslatable(f'data point {name} is not in the delivery specification')
        scored = sorted(((lcp(c, context), lcp(c, target), c) for c in cands), reverse=True)
        best = scored[0][:2]
        return sorted(c for s1, s2, c in scored if (s1, s2) == best)

    def container(self, name, near):
        suffix = '/' + name.strip('/')
        cands = [p for p in self.containers if p.endswith(suffix) or p == name]
        if not cands:
            raise Untranslatable(f'container {name} is not in the delivery specification')
        scored = sorted(((lcp(c, near), -len(c), c) for c in cands), reverse=True)
        if len(scored) > 1 and scored[0][:2] == scored[1][:2]:
            raise Untranslatable(f'container {name} is ambiguous')
        return scored[0][2]

    def repeating_ancestor(self, path, floor):
        """Deepest container on `path` (at or below `floor` depth) that can repeat."""
        parts = path.split('/')
        for i in range(len(parts), floor, -1):
            p = '/'.join(parts[:i])
            c = self.containers.get(p)
            if c and (c['max'] is None and c['min'] is not None or (c['max'] or 0) > 1):
                return p
        return None


# ---------------------------------------------------------------- H-1 rule logic parser

TOKEN = re.compile(r'\s*(?:(?P<str>"[^"]*")|(?P<num>\d[\d,]*(?:\.\d+)?)(?![\w-])|(?P<op><>|<=|>=|=|<|>)|'
                   r'(?P<p>[(),])|(?P<word>[@A-Za-z_][\w:/\-.\']*))')


# Evident typos in H-1 rule logic, corrected so the rule can run. Each names something that does not exist
# in the delivery specification next to the one that does. Recorded in rules.json under "corrections".
TYPOS = [
    (r'\bIMRPOVEMENT\b', 'IMPROVEMENT', 'IMRPOVEMENT'),
    (r'\bConstructionMethod\b(?!Type)', 'ConstructionMethodType', 'ConstructionMethod'),
    (r'\bin in\b', 'in', 'in in'),
]
CORRECTED = defaultdict(list)


def normalize(text, rule_id=None):
    t = clean(text).replace('\u201c', '"').replace('\u201d', '"').replace('\u2019', "'")
    for pat, rep, label in TYPOS:
        if re.search(pat, t):
            t = re.sub(pat, rep, t)
            if rule_id:
                CORRECTED[rule_id].append(f'"{label}" read as "{rep}"')
    # Phrasings with one reading, rewritten into the grammar below.
    t = t.replace('is not a valid 2-character US State or Territory Code', 'is not a valid state code')
    t = re.sub(r'\s+for a given comp\b', '', t)
    t = re.sub(r"'([A-Za-z]+)'", r'"\1"', t)          # 'Dwelling' -> "Dwelling"
    t = t.replace('""', '"').replace('"\'', '"')
    t = re.sub(r'\s+', ' ', t).strip().rstrip('.').strip()
    t = re.sub(r'\b(AND|And)\b', 'and', t)
    t = re.sub(r'\b(OR|Or)\b', 'or', t)
    t = re.sub(r'^(IF|if)\b', 'If', t)
    t = re.sub(r'\bis not = ', '<> ', t)
    t = re.sub(r'\bnot = ', '<> ', t)
    t = re.sub(r'\bdoes not equal\b', '<>', t)
    return t


def tokenize(t):
    out, pos = [], 0
    while pos < len(t):
        m = TOKEN.match(t, pos)
        if not m or m.end() == pos:
            raise Untranslatable(f'cannot tokenize near {t[pos:pos + 30]!r}')
        pos = m.end()
        kind = m.lastgroup
        val = m.group(kind)
        if kind == 'str':
            val = val[1:-1]
        elif kind == 'num':
            val = float(val.replace(',', ''))
        out.append((kind, val))
    return out


class Parser:
    def __init__(self, tokens, res, target_path, target_name):
        self.t = tokens
        self.i = 0
        self.res = res
        self.target_path = target_path      # full container path of the rule's primary data point
        self.target_name = target_name
        self.scopes = []
        self.refs = []                      # (node, name, context) resolved once the anchor is known

    # token helpers
    def peek(self, k=0):
        j = self.i + k
        return self.t[j] if j < len(self.t) else (None, None)

    def word(self, k=0):
        kind, val = self.peek(k)
        return val.lower() if kind == 'word' else None

    def accept_words(self, *words):
        for k, w in enumerate(words):
            if self.word(k) != w:
                return False
        self.i += len(words)
        return True

    def expect(self, kind, val=None):
        k, v = self.peek()
        if k != kind or (val is not None and v != val):
            raise Untranslatable(f'expected {val or kind} at token {self.i}, got {v!r}')
        self.i += 1
        return v

    def is_name(self, k=0):
        kind, val = self.peek(k)
        return kind == 'word' and (val.startswith('@') or re.fullmatch(r'[A-Z][A-Za-z0-9]+', val) is not None) \
            and not val.isupper()

    def is_container(self, k=0):
        kind, val = self.peek(k)
        return kind == 'word' and re.fullmatch(r'[A-Z][A-Z0-9_]*(?:/[A-Z][A-Z0-9_]*)*', val) is not None and '_' in val + '_' \
            and val.upper() == val

    # grammar
    def rule(self):
        if not self.accept_words('if'):
            raise Untranslatable('rule logic does not start with "If"')
        node = self.top()
        while self.peek()[0] is not None and self.scope():
            pass
        if self.peek()[0] is not None:
            raise Untranslatable(f'unparsed text from token {self.i}: {self.peek()[1]!r}')
        return node

    # Outside parentheses, H-1 writes "If A or B, and C is not provided" meaning (A or B) and C:
    # at the top level "or" binds tighter than "and". Inside parentheses the usual precedence applies.
    def top(self):
        items = [self.top_disj()]
        while True:
            save = self.i
            if self.peek() == ('p', ','):
                self.i += 1
            if self.accept_words('and'):
                items.append(self.top_disj())
            else:
                self.i = save
                break
        return items[0] if len(items) == 1 else ['and'] + items

    def top_disj(self):
        items = [self.unary()]
        while True:
            save = self.i
            if self.peek() == ('p', ','):
                self.i += 1
            if self.accept_words('or'):
                items.append(self.unary())
            else:
                self.i = save
                break
        return items[0] if len(items) == 1 else ['or'] + items

    def expr(self):
        items = [self.conj()]
        while True:
            save = self.i
            if self.peek() == ('p', ','):
                self.i += 1
            if self.accept_words('or'):
                items.append(self.conj())
            else:
                self.i = save
                break
        return items[0] if len(items) == 1 else ['or'] + items

    def conj(self):
        items = [self.unary()]
        while True:
            save = self.i
            if self.peek() == ('p', ','):
                self.i += 1
            if self.accept_words('and'):
                items.append(self.unary())
            else:
                self.i = save
                break
        return items[0] if len(items) == 1 else ['and'] + items

    def unary(self):
        if self.peek() == ('p', '('):
            self.i += 1
            node = self.expr()
            self.expect('p', ')')
            self.scope()
            return node
        return self.atom()

    def with_clause(self, cpath):
        """'with' Name op value ['and' Name op value ...], names resolved inside container cpath."""
        if not self.accept_words('with'):
            return None
        conds = [self.comparison(context=cpath, inner=True)]
        while self.word() == 'and' and self.is_name(1) and self.peek(2)[0] == 'op':
            try:
                self.res.leaf(self.peek(1)[1], cpath, cpath)
            except Untranslatable:
                break
            if not self.res.leaf(self.peek(1)[1], cpath, cpath)[0].startswith(cpath + '/'):
                break
            self.i += 1
            conds.append(self.comparison(context=cpath, inner=True))
        return conds[0] if len(conds) == 1 else ['and'] + conds

    def container_ref(self):
        if not self.is_container():
            raise Untranslatable(f'expected a container name at token {self.i}: {self.peek()[1]!r}')
        name = self.peek()[1]
        self.i += 1
        return self.res.container(name, self.target_path)

    def count_atom(self, op, n):
        cpath = self.container_ref()
        cond = self.with_clause(cpath)
        return ['count', cpath, cond, op, n]

    def atom(self):
        # instance counts
        if self.accept_words('there', 'is', 'no', 'instance', 'or', 'more', 'than', 'one', 'instance', 'of'):
            return self.count_atom('!=', 1)
        if self.accept_words('there', 'is', 'more', 'than', 'one', 'instance', 'of') or \
                self.accept_words('there', 'are', 'more', 'than', 'one', 'instance', 'of'):
            return self.count_atom('>', 1)
        if self.accept_words('there', 'is', 'no', 'instance', 'of') or self.accept_words('there', 'are', 'no', 'instances', 'of'):
            return self.count_atom('==', 0)
        if self.accept_words('no', 'instances', 'of') or self.accept_words('no', 'instance', 'of'):
            node = self.count_atom('==', 0)
            if not (self.accept_words('are', 'provided') or self.accept_words('is', 'provided')):
                raise Untranslatable('expected "are provided"')
            return node
        if self.accept_words('more', 'than', 'one', 'instance', 'of'):
            cpath = self.container_ref()
            if not self.accept_words('has'):
                raise Untranslatable('expected "has"')
            cond = self.comparison(context=cpath, inner=True)
            return ['count', cpath, cond, '>', 1]
        if self.accept_words('an', 'instance', 'of'):
            # "an instance of C [with X = v] is not provided" / "an instance of C is not provided with X = v"
            cpath = self.container_ref()
            cond = self.with_clause(cpath)
            if not self.accept_words('is', 'not', 'provided'):
                raise Untranslatable('expected "is not provided"')
            if cond is None:
                cond = self.with_clause(cpath)
            return ['count', cpath, cond, '==', 0]

        if not self.is_name():
            raise Untranslatable(f'expected a data point name at token {self.i}: {self.peek()[1]!r}')
        name = self.peek()[1]
        self.i += 1
        if self.accept_words('is', 'not', 'provided'):
            node = ['missing', self.ref(name)]
            self.scope()
            return node
        if self.accept_words('is', 'provided'):
            node = ['present', self.ref(name)]
            self.scope()
            return node
        if self.accept_words('is', 'not', 'unique', 'across', 'all', 'instances', 'of'):
            cpath = self.container_ref()
            cond = self.with_clause(cpath)
            node = ['unique', self.ref(name, context=cpath), cpath, cond]
            self.scope()
            return node
        if self.accept_words('does', 'not', 'include', 'year'):
            # "...year, month and day (YYYY-MM-DD format)" / "...year and month (YYYY-MM format)"
            start = self.i
            while self.peek()[0] is not None and not (self.word() in ('in', 'for') and self.word(1) in ('a', 'the')):
                self.i += 1
            words = [str(v).lower() for k, v in self.t[start:self.i] if k == 'word']
            fmt = 'YYYY-MM-DD' if 'day' in words else ('YYYY-MM' if 'month' in words else None)
            if fmt is None:
                raise Untranslatable('unrecognised date format rule')
            node = ['datefmt', self.ref(name), fmt]
            self.scope()
            return node
        if self.accept_words('is', 'in', 'the', 'future'):
            return ['future', self.ref(name)]
        if self.accept_words('is', 'more', 'than'):
            kind, n = self.peek()
            if kind != 'num':
                raise Untranslatable('expected a day count')
            self.i += 1
            if not self.accept_words('days', 'before', 'the', 'current', 'date'):
                raise Untranslatable('expected "days before the current date"')
            return ['older', self.ref(name), int(n)]
        if self.accept_words('is', 'not', 'in', 'the', 'required', 'format'):
            if name != 'PostalCode':
                raise Untranslatable('format rule for a data point other than PostalCode')
            self.skip_parenthetical()
            return ['zip', self.ref(name)]
        if self.accept_words('is', 'not', 'a', 'valid', 'state', 'code'):
            return ['state', self.ref(name)]
        self.i -= 1
        node = self.comparison(context=None)
        self.scope()
        return node

    def skip_parenthetical(self):
        if self.peek() == ('p', '('):
            depth = 0
            while self.peek()[0] is not None:
                if self.peek() == ('p', '('):
                    depth += 1
                elif self.peek() == ('p', ')'):
                    depth -= 1
                self.i += 1
                if depth == 0:
                    break

    def comparison(self, context, inner=False):
        if not self.is_name():
            raise Untranslatable(f'expected a data point name at token {self.i}: {self.peek()[1]!r}')
        name = self.peek()[1]
        self.i += 1
        if self.accept_words('is', 'not'):
            op = '<>'
        elif self.accept_words('is'):
            op = '='
        else:
            kind, op = self.peek()
            if kind != 'op':
                raise Untranslatable(f'expected an operator after {name}, got {op!r}')
            self.i += 1
        lhs = self.ref(name, context=context)
        rhs = self.value(context, inner)
        if isinstance(rhs, list) and rhs and rhs[0] == 'countof':
            if op not in ('<>', '='):
                raise Untranslatable('count comparison with an unsupported operator')
            return ['counteq', op, lhs, rhs[1], rhs[2]]
        return ['cmp', op, lhs, rhs]

    def value(self, context, inner):
        kind, val = self.peek()
        if kind == 'str':
            self.i += 1
            alts = [val]
            while self.word() == 'or' and self.peek(1)[0] == 'str':
                alts.append(self.peek(1)[1])
                self.i += 2
            return {'str': alts}
        if kind == 'p' and val == '(' and self.peek(1)[0] == 'str':
            self.i += 1
            alts = [self.expect('str')]
            while self.accept_words('or'):
                alts.append(self.expect('str'))
            self.expect('p', ')')
            return {'str': alts}
        if kind == 'num':
            self.i += 1
            return {'num': val}
        if self.accept_words('the', 'number', 'of', 'instances', 'of'):
            cpath = self.container_ref()
            return ['countof', cpath, self.with_clause(cpath)]
        if kind == 'word' and val in ('true', 'false'):
            self.i += 1
            return {'str': [val]}
        if self.is_name() and not inner:
            self.i += 1
            return self.ref(val, context=context)
        raise Untranslatable(f'unsupported value at token {self.i}: {val!r}')

    def scope(self):
        """Optional trailing scope; records it and returns True when one was consumed."""
        start = self.i
        if self.accept_words('in', 'a', 'given', 'instance', 'of') or self.accept_words('in', 'the', 'given', 'instance', 'of') \
                or self.accept_words('for', 'a', 'given', 'instance', 'of'):
            cpath = self.container_ref()
            self.scopes.append(('instance', cpath, self.with_clause(cpath)))
            return True
        if self.accept_words('for', 'the', 'instance', 'of') or self.accept_words('in', 'the', 'instance', 'of'):
            cpath = self.container_ref()
            cond = self.with_clause(cpath)
            if cond is None:
                raise Untranslatable('instance scope without a "with" condition')
            self.scopes.append(('instance', cpath, cond))
            return True
        if self.accept_words('for', 'one', 'or', 'more', 'instances', 'of'):
            cpath = self.container_ref()   # "@ValuationUseType is not provided for one or more instances of PROPERTY"
            self.scopes.append(('instance', cpath, None))
            return True
        if self.word() == 'in' and self.is_container(1):
            # "is not provided in ROOF_DETAIL [in a given instance of ...]": names the container the data point
            # lives in, which its path already says. "in SERVICE container with X = v" filters that container.
            cname = self.peek(1)[1]
            self.i += 2
            if self.accept_words('container'):
                cpath = self.res.container(cname, self.target_path)
                cond = self.with_clause(cpath)
                if cond is None:
                    return self._fail(start)
                self.scopes.append(('instance', cpath, cond))
                return True
            if self.word() in ('for', 'in'):
                return self.scope() or self._fail(start)
            return True
        if self.word() == 'for' and self.is_name(1) and self.peek(2)[0] == 'op':
            self.i += 1
            cond = self.comparison(context=None, inner=True)
            self.scopes.append(('filter', cond))
            return True
        return False

    def _fail(self, start):
        self.i = start
        return False

    def ref(self, name, context=None):
        node = {'ref': name}
        self.refs.append((node, name, context))
        return node


def walk_refs(node, fn):
    if isinstance(node, dict):
        if 'ref' in node:
            fn(node)
    elif isinstance(node, list):
        for x in node:
            walk_refs(x, fn)


def translate(rule, res):
    logic = normalize(rule['Rule Logic'], clean(rule['Message ID']))
    if not logic:
        raise Untranslatable('no rule logic')
    if re.match(r'(?i)^for each\b', logic):
        raise Untranslatable('"for each combination" rules compare rows across properties; not translated')
    if 'xlink:arcrole' in logic or 'RELATIONSHIP' in logic:
        raise Untranslatable('RELATIONSHIP link rules need the xlink graph; not translated')
    if re.search(r'\byear of\b|\bYYYY-MM of\b|plus\b|\bsum of\b|\bwithin 1 year\b|\bat least one instance\b|'
                 r'\bany of\b|\beither\b|\bfor each role\b|\bsecond instance\b|\bone or more additional\b|\band/or\b', logic):
        raise Untranslatable('rule needs date arithmetic, sums or nested quantifiers; not translated')

    target_name = re.sub(r'\s*\(.*\)$', '', clean(rule['Data Point Name / Value']))
    xpath = clean(rule['xPath']).lstrip('.').strip('/')
    if not xpath:
        raise Untranslatable('no xPath')
    tpath = res.container(xpath, VA)

    p = Parser(tokenize(logic), res, tpath, target_name)
    expr = p.rule()

    # ---- anchor
    affected = set()
    for part in clean(rule['Property Affected']).split(';'):
        key = re.sub(r'\s*#n$', '', part.strip().lower())
        if key in AFFECTED:
            affected.add(AFFECTED[key])
    instance_scopes = {}
    for s in p.scopes:
        if s[0] == 'instance':
            instance_scopes.setdefault((s[1], json.dumps(s[2], sort_keys=True)), s)
    filters = [s[1] for s in p.scopes if s[0] == 'filter']
    if len(instance_scopes) > 1:
        raise Untranslatable('more than one instance scope')
    if instance_scopes and filters:
        raise Untranslatable('instance scope combined with a "for" filter')
    if instance_scopes:
        _, apath, acond = next(iter(instance_scopes.values()))
    elif filters:
        f = filters[0]
        if f[0] != 'cmp' or 'ref' not in f[2]:
            raise Untranslatable('unsupported "for" filter')
        fpath = res.leaf(f[2]['ref'], tpath, tpath)[0]
        if f[2]['ref'] == '@ValuationUseType':
            apath = PROPERTY
        else:
            common = '/'.join(fpath.split('/')[:lcp(fpath, tpath)])
            apath = res.repeating_ancestor(common, 1)
            if not apath:
                raise Untranslatable('no repeating container for the "for" filter')
        acond = f
        p.refs.append((f[2], f[2]['ref'], apath))
    elif tpath == PROPERTY or tpath.startswith(PROPERTY + '/'):
        apath, acond = PROPERTY, None
    else:
        apath, acond = 'MESSAGE', None

    # ---- resolve every name to full paths, nearest the context it is evaluated in
    for node, name, context in p.refs:
        ctx = context or apath
        node['paths'] = res.leaf(name, ctx, tpath if name == target_name or context is None else ctx)

    if target_name and not re.search(r'\b' + re.escape(target_name) + r'\b', logic) and target_name.split('/')[0] not in logic:
        pass  # the primary data point is not always named in the logic (e.g. count rules); that is fine

    return {
        'id': clean(rule['Message ID']), 'uid': clean(rule['Unique ID']), 'sev': clean(rule['Severity']),
        'msg': clean(rule['Message Text']), 'affected': sorted(affected),
        'anchor': {'path': apath, 'with': acond},
        'target': {'path': tpath, 'name': target_name},
        'expr': expr,
    }


def public_reason(reason):
    """Group the translator's reasons into the categories shown on the website."""
    if reason.startswith('RELATIONSHIP'):
        return 'Links between parts of the report (RELATIONSHIP / xlink)'
    if reason.startswith('"for each combination"'):
        return 'Row-by-row comparison across all comparables'
    if reason.startswith('rule needs'):
        return 'Date arithmetic, sums or nested conditions'
    if 'not in the delivery specification' in reason or 'ambiguous' in reason:
        return 'Names a data point or container the delivery specification does not define at that location'
    return 'Wording without one unambiguous machine reading'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--spec-dir', required=True, help='folder holding a1.xlsx and h1.xlsx (GSE appendices)')
    ap.add_argument('--a1', default='a1.xlsx')
    ap.add_argument('--h1', default='h1.xlsx')
    args = ap.parse_args()
    a1_file, h1_file = Path(args.spec_dir) / args.a1, Path(args.spec_dir) / args.h1

    a1 = openpyxl.load_workbook(a1_file, read_only=True)
    spec_sheet = next(n for n in a1.sheetnames if n.startswith('UAD Delivery Spec') and 'Marked' not in n)
    containers, datapoints = build_spec(sheet_rows(a1, spec_sheet))
    report_ids = sorted({e for k, v in datapoints.items() if k.endswith('/ValuationReportContentIdentifier') for x in v for e in x['enum']})

    h1 = openpyxl.load_workbook(h1_file, read_only=True)
    rules_sheet = next(n for n in h1.sheetnames if n.startswith('UAD Compliance Rules') and 'Marked' not in n)
    h1_rows = [r for r in sheet_rows(h1, rules_sheet) if clean(r.get('Message ID'))]

    res = Resolver(containers, datapoints)
    rules, skipped = [], []
    for r in h1_rows:
        try:
            rules.append(translate(r, res))
        except Untranslatable as exc:
            skipped.append({'id': clean(r['Message ID']), 'uid': clean(r['Unique ID']), 'sev': clean(r['Severity']),
                            'msg': clean(r['Message Text']), 'reason': str(exc)})

    out = {
        'generatedBy': 'scripts/uad/build_rules.py',
        'sources': {
            'a1': {'title': 'Appendix A-1: URAR Delivery Specification', 'sheet': spec_sheet, 'sha256': sha256(a1_file)},
            'h1': {'title': 'Appendix H-1: URAR Compliance Rules', 'sheet': rules_sheet, 'sha256': sha256(h1_file)},
        },
        'mismoVersion': '3.6.0366',
        'reportTypes': report_ids,
        'containers': containers,
        'datapoints': datapoints,
        'rules': rules,
        'notImplemented': skipped,
        'corrections': {k: v for k, v in sorted(CORRECTED.items()) if k in {r['id'] for r in rules}},
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, separators=(',', ':'), sort_keys=True, ensure_ascii=False))

    # Small summary for the website (product page and docs state exact coverage from this).
    sev = lambda xs: dict(sorted(Counter('Fatal' if x['sev'].startswith('Fatal') else 'Warning' for x in xs).items()))
    summary = {
        'sources': {k: {'title': v['title'], 'sheet': v['sheet']} for k, v in out['sources'].items()},
        'rulesTotal': len(h1_rows),
        'rulesImplemented': len(rules),
        'implementedBySeverity': sev(rules),
        'notImplemented': [{'id': x['id'], 'severity': 'Fatal' if x['sev'].startswith('Fatal') else 'Warning',
                            'reason': public_reason(x['reason'])} for x in sorted(skipped, key=lambda x: x['id'])],
        'corrections': out['corrections'],
        'dataPoints': len(datapoints),
    }
    COVERAGE.write_text(json.dumps(summary, indent=1) + '\n')
    n = len(h1_rows)
    print(f'{len(rules)}/{n} H-1 rules translated ({len(rules) * 100 // n}%), {len(skipped)} not implemented; '
          f'{len(datapoints)} data points, {len(containers)} containers -> {OUT.relative_to(ROOT)} ({OUT.stat().st_size:,} bytes)')


if __name__ == '__main__':
    sys.exit(main())

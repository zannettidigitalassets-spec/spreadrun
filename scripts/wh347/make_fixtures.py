"""Make the synthetic WH-347 test payrolls: a fictional contractor, one week, invented workers and invented rates.

    python3 scripts/wh347/make_fixtures.py

Writes to pylib/tests/fixtures/ and public/samples/:
  wh347-pass.xlsx / wh347-pass.json   a payroll that passes every check (also the template users can copy)
  wh347-fail.xlsx / wh347-fail.json   the same week with one planted problem of each kind
  wh347-*.csv                         the three tables of the pass payroll as plain CSV
The wage rates are made up and come from no real wage determination. Needs openpyxl (build time only).
"""
import copy
import csv
import io
import json
import shutil
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'pylib' / 'spreadrun_api' / 'validators' / 'wh347'))
import engine  # noqa: E402

FIX = ROOT / 'pylib' / 'tests' / 'fixtures'
PUB = ROOT / 'public' / 'samples'
D = Decimal
C = lambda x: D(x).quantize(D('0.01'), rounding=ROUND_HALF_UP)  # noqa: E731

HEADER = {
    'project_name': 'Example Community Center Renovation',
    'project_no': 'EX-2026-0147',
    'payroll_no': '6',
    'contractor_name': 'Example Builders LLC (fictional)',
    'project_location': 'Example County, Example State',
    'wage_determination_no': 'EX20260001 Mod 0 (invented rates)',
    'week_ending': '2026-09-26',
    'contractor_address': '100 Sample Way, Example City',
    'cwhssa': 'yes',
}
WD = [  # classification, base, fringe: invented
    ('Electrician', '38.50', '18.25'),
    ('Carpenter', '31.20', '14.10'),
    ('Laborer: Common or General', '24.15', '11.40'),
    ('Operator: Backhoe/Excavator', '35.80', '16.05'),
]
APPR = [  # classification, level, wage %, fringe % (blank = program silent: full fringe), ratio, program
    ('Electrician', '2', '60', '', '1:1', 'Example Electrical JATC (fictional)'),
    ('Electrician', '3', '70', '', '1:1', 'Example Electrical JATC (fictional)'),
]
COLS = (['entry_no', 'last_name', 'first_name', 'middle_initial', 'worker_id', 'worker_type', 'apprentice_level',
         'classification'] + [f'st_{d}' for d in range(1, 8)] + [f'ot_{d}' for d in range(1, 8)]
        + ['total_hours', 'st_rate', 'ot_rate', 'fringe_credit_hourly', 'fringe_credit', 'cash_in_lieu_hourly',
           'cash_in_lieu', 'gross_project', 'gross_all_work', 'tax_withholding', 'fica', 'other_deductions',
           'total_deductions', 'net_pay'])


def row(entry, last, wid, wtype, level, cls, st, ot, st_rate, ot_rate, fb_hr='0', cash_hr='0', weekly=True,
        other_work='0', other_ded='0'):
    st, ot = [D(x) for x in st], [D(x) for x in ot]
    hrs = sum(st) + sum(ot)
    st_rate, ot_rate = D(st_rate), D(ot_rate) if ot_rate else None
    fb, cash = C(D(fb_hr) * hrs), C(D(cash_hr) * hrs)
    gross = C(st_rate * sum(st) + (ot_rate or 0) * sum(ot) + cash)
    r = dict(zip(COLS, [''] * len(COLS)))
    r.update({'entry_no': str(entry), 'last_name': last, 'first_name': 'Worker', 'middle_initial': '',
              'worker_id': wid, 'worker_type': wtype, 'apprentice_level': level, 'classification': cls,
              'total_hours': str(hrs), 'st_rate': str(st_rate), 'ot_rate': str(ot_rate) if ot_rate else '',
              'fringe_credit_hourly': fb_hr if fb_hr != '0' else '', 'fringe_credit': str(fb),
              'cash_in_lieu_hourly': cash_hr if cash_hr != '0' else '', 'cash_in_lieu': str(cash),
              'gross_project': str(gross)})
    for d in range(7):
        r[f'st_{d + 1}'] = str(st[d]) if st[d] else ''
        r[f'ot_{d + 1}'] = str(ot[d]) if ot[d] else ''
    r['_gross'] = gross
    r['_weekly'] = weekly
    r['_other_work'] = D(other_work)
    r['_other_ded'] = D(other_ded)
    return r


def finish(rows):
    """Fill 7B, 8 and 9 once per worker, on the worker's first row."""
    by_entry = {}
    for r in rows:
        by_entry.setdefault(r['entry_no'], []).append(r)
    for entry, rs in by_entry.items():
        g_all = sum(r['_gross'] for r in rs) + rs[0]['_other_work']
        tax, fica = C(g_all * D('0.10')), C(g_all * D('0.0765'))
        other = rs[0]['_other_ded']
        total = tax + fica + other
        rs[0].update({'gross_all_work': str(g_all), 'tax_withholding': str(tax), 'fica': str(fica),
                      'other_deductions': str(other), 'total_deductions': str(total), 'net_pay': str(g_all - total)})
    return rows


W5 = ['8', '8', '8', '8', '8', '0', '0']
Z = ['0'] * 7


def pass_rows():
    return finish([
        # Electrician, plan credit for the full fringe, 4 hours overtime on Saturday at 1.5 times the rate paid.
        row(1, 'Sample-01', '4410', 'J', '', 'Electrician', W5, ['0', '0', '0', '0', '0', '4', '0'], '39.00', '58.50',
            fb_hr='18.25'),
        # Registered apprentice, level 2 (60 percent), same plan, one journeyworker electrician on site each day.
        row(2, 'Sample-02', '7782', 'RA', '2', 'Electrician', W5, Z, '23.10', '', fb_hr='18.25'),
        # Carpenter paid the fringe in cash in lieu of benefits.
        row(3, 'Sample-03', '1209', 'J', '', 'Carpenter', W5, Z, '31.20', '', cash_hr='14.10'),
        # Worker in two classifications in the same week, part plan and part cash, with other non-project work.
        row(4, 'Sample-04', '5530', 'J', '', 'Laborer: Common or General', ['8', '8', '8', '0', '0', '0', '0'], Z,
            '24.15', '', fb_hr='6.00', cash_hr='5.40', other_work='120.00', other_ded='15.00'),
        row(4, 'Sample-04', '5530', 'J', '', 'Operator: Backhoe/Excavator', ['0', '0', '0', '8', '8', '0', '0'], Z,
            '35.80', '', fb_hr='6.00', cash_hr='10.05'),
    ])


def fail_rows():
    rows = pass_rows()
    rows = [copy.deepcopy(r) for r in rows]
    rows[0]['ot_rate'] = '55.00'                     # below 1.5 x 38.50 = 57.75: WH-RATE-OT-BASE
    rows[0]['net_pay'] = str(D(rows[0]['net_pay']) + 25)  # WH-NET-PAY
    rows[2]['st_rate'] = '30.00'                     # below 31.20: WH-RATE-BASE (and fringe short)
    rows[2]['worker_id'] = '123-45-6789'             # full SSN: WH-ID-FULL-SSN
    rows[3]['total_hours'] = '25'                    # column 4 says 24: WH-HOURS-TOTAL
    rows[3]['total_deductions'] = str(D(rows[3]['total_deductions']) + 1)  # WH-DEDUCTIONS-TOTAL (and net pay)
    rows[4]['classification'] = 'Operator: Crane'    # not on the WD: WH-WD-CLASSIFICATION
    extra = finish([
        # 44 straight time hours, no overtime: WH-OT-HOURS. Plan credit below the fringe rate: WH-FRINGE-SHORT.
        row(5, 'Sample-05', '6601', 'J', '', 'Laborer: Common or General', ['9', '9', '9', '9', '8', '0', '0'], Z,
            '24.15', '', fb_hr='9.00'),
        # A second electrical apprentice the same days as the first: two apprentices, one journeyworker (ratio 1:1).
        row(6, 'Sample-06', '3318', 'RA', '3', 'Electrician', W5, Z, '26.95', '', fb_hr='18.25'),
        # Apprentice level with no program wage schedule supplied: WH-APPR-PROGRAM.
        row(7, 'Sample-07', '9920', 'RA', '5', 'Electrician', ['0', '0', '0', '0', '0', '4', '0'], Z, '30.00', '',
            fb_hr='18.25'),
    ])
    rows[2]['gross_project'] = str(D(rows[2]['gross_project']) + D('7.77'))  # WH-GROSS-MISMATCH (warning)
    return rows + extra


def csv_text(cols, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator='\n')
    w.writerow(cols)
    for r in rows:
        w.writerow([r[c] for c in cols] if isinstance(r, dict) else r)
    return buf.getvalue()


def as_json(rows):
    return json.dumps({
        'header': {k: v for k, v in HEADER.items() if k != 'cwhssa'},
        'cwhssa': True,
        'payrollCsv': csv_text(COLS, rows),
        'wageDeterminationCsv': csv_text(['classification', 'base_rate', 'fringe_rate'], WD),
        'apprenticeshipCsv': csv_text(['classification', 'level', 'wage_percent', 'fringe_percent', 'ratio',
                                       'program_name'], APPR),
    }, indent=2)


def number(v):
    try:
        return float(v) if '.' in v else int(v)
    except ValueError:
        return v


def as_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Header'
    ws.append(['field', 'value'])
    for k, v in HEADER.items():
        ws.append([k, v])
    ws = wb.create_sheet('Payroll')
    ws.append(COLS)
    for r in rows:
        ws.append([number(r[c]) if c not in ('worker_id', 'entry_no', 'apprentice_level') else r[c] for c in COLS])
    ws = wb.create_sheet('Wage Determination')
    ws.append(['classification', 'base_rate', 'fringe_rate'])
    for c, b, f in WD:
        ws.append([c, float(b), float(f)])
    ws = wb.create_sheet('Apprenticeship')
    ws.append(['classification', 'level', 'wage_percent', 'fringe_percent', 'ratio', 'program_name'])
    for r in APPR:
        ws.append(list(r))
    for s in wb.worksheets:
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width = max(10, min(36, max(len(str(c.value or '')) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    FIX.mkdir(parents=True, exist_ok=True)
    PUB.mkdir(parents=True, exist_ok=True)
    out = {}
    for name, rows in (('pass', pass_rows()), ('fail', fail_rows())):
        out[f'wh347-{name}.json'] = as_json(rows).encode()
        out[f'wh347-{name}.xlsx'] = as_xlsx(rows)
    p = pass_rows()
    out['wh347-payroll.csv'] = csv_text(COLS, p).encode()
    out['wh347-wage-determination.csv'] = csv_text(['classification', 'base_rate', 'fringe_rate'], WD).encode()
    out['wh347-apprenticeship.csv'] = csv_text(['classification', 'level', 'wage_percent', 'fringe_percent', 'ratio',
                                                'program_name'], APPR).encode()
    for n in ('wh347-pass.json', 'wh347-pass.xlsx'):
        r = engine.validate(out[n])
        if r['status'] != 'PASS':
            raise SystemExit(f'{n} does not pass: {r["ruleCounts"]}')
    want = {'WH-RATE-OT-BASE', 'WH-NET-PAY', 'WH-RATE-BASE', 'WH-ID-FULL-SSN', 'WH-HOURS-TOTAL', 'WH-DEDUCTIONS-TOTAL',
            'WH-WD-CLASSIFICATION', 'WH-OT-HOURS', 'WH-FRINGE-SHORT', 'WH-APPR-RATIO', 'WH-APPR-PROGRAM',
            'WH-GROSS-MISMATCH'}
    for n in ('wh347-fail.json', 'wh347-fail.xlsx'):
        r = engine.validate(out[n])
        if r['status'] != 'FAIL' or not want <= set(r['ruleCounts']):
            raise SystemExit(f'{n} missing {want - set(r["ruleCounts"])}: {r["ruleCounts"]}')
    for n, b in out.items():
        (FIX / n).write_bytes(b)
        shutil.copy(FIX / n, PUB / n)
    r = engine.validate(out['wh347-fail.json'])
    print(f'wh347 samples written; fail: {r["findingCount"]} findings {sorted(r["ruleCounts"])}')


if __name__ == '__main__':
    main()

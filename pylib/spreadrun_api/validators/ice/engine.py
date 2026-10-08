"""SpreadRun ICE incurred cost submission adequacy pre-check.

Checks a contractor's annual incurred cost submission workbook for adequacy, the way DCAA's checklist frames it:
are the required schedules there, does the math foot, do the schedules tie to each other, is the certificate there,
and is it on time. SpreadRun's own engine (not a DataForge copy). No network calls.

It never decides whether a cost is allowable (FAR Part 31). That is judgment, not math, and the report says so.

Input: one .xlsx workbook (any layout: the DCAA ICE model is not required), with the fiscal year end and, optionally,
the date to measure the deadline from, as query parameters.

Rules, read October 7, 2026:
  FAR 52.216-7(d)(2)(i)      adequate final indirect cost rate proposal due within 6 months after fiscal year end;
                             extensions only in writing, for exceptional circumstances.
  FAR 52.216-7(d)(2)(iii)    the content of an adequate proposal, items (A) through (O): Schedules A to O.
  FAR 42.703-2, 52.242-4     certificate of final indirect costs; without it no proposal is accepted, and the
                             contracting officer may set rates unilaterally. Signed at vice president or CFO level.
  FAR 42.705-1(b)(1)         the auditor reviews adequacy and describes inadequacies in writing.
  DCAA Checklist for Determining Adequacy of Contractor Incurred Cost Proposal, Version 3.4 (December 2021): 47
                             items by schedule, plus the instruction that math and formulas in each schedule be
                             accurate; contractor reports may stand in for the example schedules.

Every amount the report shows is read from, or recomputed from, the workbook. Text from the workbook (names,
addresses, contract numbers, account names) is never repeated: findings name the tab and cell range.
"""
import datetime as dt
import hashlib
import io
import re
import time
import zipfile
from collections import Counter, defaultdict

MAX_BYTES = 4_000_000
TIME_BUDGET_S = 20.0
MAX_ROWS = 20000
MAX_COLS = 150
MAX_FINDINGS = 400
TOL = 1.0              # dollars: ICE schedules are in whole dollars, so sums may differ by rounding of a dollar
RATE_TOL_FLOOR = 1e-9

SOURCES = {
    'far52.216-7': 'FAR 52.216-7(d)(2), Allowable Cost and Payment: the 6-month deadline and items (A) to (O) of an '
                   'adequate final indirect cost rate proposal',
    'far42.703-2': 'FAR 42.703-2, Certificate of indirect costs (and the clause at FAR 52.242-4)',
    'far42.705-1': 'FAR 42.705-1(b)(1), adequacy review of the final indirect cost rate proposal',
    'dcaa-checklist': 'DCAA Checklist for Determining Adequacy of Contractor Incurred Cost Proposal, Version 3.4 '
                      '(December 2021)',
}

# Schedule letter -> (short name, title words that identify it, FAR 52.216-7(d)(2)(iii) item)
SCHEDULES = {
    'A': ('Summary of claimed indirect rates', r'SUMMARY OF (ALL )?CLAIMED INDIRECT|CLAIMED INDIRECT (EXPENSE )?RATES|INDIRECT RATES SUMMARY'),
    'B': ('General and administrative expenses', r'GENERAL\s*(AND|&)\s*ADMINISTRATIVE|G\s*&\s*A EXPENSE'),
    'C': ('Overhead expenses', r'OVERHEAD'),
    'D': ('Occupancy or intermediate pool expenses', r'OCCUPANCY|INTERMEDIATE (INDIRECT|ALLOCATION|POOL)'),
    'E': ('Claimed allocation bases', r'ALLOCATION BASES?'),
    'F': ('Facilities capital cost of money', r'COST OF MONEY|FACILITIES CAPITAL'),
    'G': ('Reconciliation of books of account to claimed direct costs', r'RECONCILIATION OF BOOKS|BOOKS OF ACCOUNT'),
    'H': ('Direct costs by contract with indirect expense applied', r'DIRECT COSTS? BY CONTRACT|SCHEDULE OF DIRECT COSTS'),
    'I': ('Cumulative costs claimed and billed', r'CUMULATIVE.*CLAIMED AND BILLED|CLAIMED AND BILLED'),
    'J': ('Subcontract information', r'SUBCONTRACT INFORMATION|LISTING OF SUBCONTRACTS'),
    'K': ('Time-and-materials and labor-hour contracts', r'TIME\s*(AND|&)\s*MATERIALS?|LABOR[\s-]HOUR'),
    'L': ('Reconciliation of IRS Form 941 payroll to labor distribution', r'941|PAYROLL'),
    'M': ('Decisions, agreements, approvals and accounting changes', r'DECISIONS|AGREEMENTS|ACCOUNTING (OR|AND|/) ?ORGANI'),
    'N': ('Certificate of final indirect costs', r'CERTIFICATE OF FINAL INDIRECT|CERTIFICATION OF FINAL INDIRECT'),
    'O': ('Contract closing information', r'CONTRACT CLOSING|PHYSICALLY COMPLETE'),
}
LETTERS = list(SCHEDULES)

# The 47 DCAA checklist items, in our own short words, with how this check covers each one.
#   (item, schedule, summary, check id or None for "review")
CHECKLIST = [
    (1, 'A', 'All claimed pools, bases and rates identified, with cost of money where it applies', 'a-rows'),
    (2, 'A', 'A cost schedule for each final pool on Schedule A (Schedules B and C)', 'present-BC'),
    (3, 'A', 'A cost schedule for each intermediate pool on Schedule A (Schedule D)', 'present-D'),
    (4, 'A', 'Pool amounts on Schedule A tie to the claimed totals on Schedules B and C', 'tie-pools'),
    (5, 'A', 'Intermediate pool bases on Schedule A tie to Schedule D', 'tie-d-base'),
    (6, 'A', 'Final pool bases on Schedule A tie to Schedule E', 'tie-e-base'),
    (7, 'B', 'G&A pool costs tie to Schedule H', None),
    (8, 'B', 'Explanatory notes for adjustments and omitted amounts', 'notes-B'),
    (9, 'B', 'Intermediate allocations appear on the source schedules', None),
    (10, 'B', 'Fringe and overhead applied to IR&D and B&P', None),
    (11, 'C', 'Overhead pool costs tie to Schedule H', None),
    (12, 'C', 'Explanatory notes for adjustments and omitted amounts', 'notes-C'),
    (13, 'C', 'Intermediate allocations appear on the source schedules', None),
    (14, 'D', 'Explanatory notes for adjustments and omitted amounts', 'notes-D'),
    (15, 'D', 'Intermediate allocations appear on the receiving schedules', None),
    (16, 'D', 'Allocation base by recipient, percentage of the base, and dollars allocated', 'd-allocation'),
    (17, 'E', 'An explanation of each base', None),
    (18, 'E', 'Base cost elements tie to the referenced schedules, with notes', None),
    (19, 'F', 'Cost of money allocation bases match the bases on Schedule A', 'tie-f-base'),
    (20, 'F', 'A separate cost of money rate for each final pool', None),
    (21, 'G', 'Direct costs per the general ledger tie to Schedule H', 'tie-g-h'),
    (22, 'G', 'Explanatory notes for adjustments and omitted amounts', 'notes-G'),
    (23, 'H', 'Flexibly priced contracts listed by contract and subtotaled by contract type', 'h-subtotals'),
    (24, 'H', 'Subcontract costs by contract tie to Schedule J', None),
    (25, 'H', 'Cost detail at the level each contract bills at', None),
    (26, 'H', 'Indirect expenses applied at the claimed rates from Schedule A', 'h-rates'),
    (27, 'H', 'Government participation calculated for each indirect pool', None),
    (28, 'H', 'Bases used for Government participation tie to Schedules E and H', None),
    (29, 'I', 'Cost detail at the level used for billing', None),
    (30, 'I', 'Fiscal year claimed amounts tie to Schedule H for cost-type contracts', 'tie-i-h'),
    (31, 'I', 'Fiscal year claimed amounts tie to Schedule K for T&M contracts', 'tie-i-k'),
    (32, 'I', 'Prior years settled costs match the prior Cumulative Allowable Cost Worksheet', None),
    (33, 'I', 'Contracts marked physically complete appear on Schedule O', 'tie-i-o'),
    (34, 'J', 'All types of subcontracts and intercompany costs included', None),
    (35, 'J', 'Full detail for each subcontract', 'j-detail'),
    (36, 'K', 'Cost detail at the level used for billing', None),
    (37, 'K', 'Every T&M or labor-hour contract on Schedule H appears on Schedule K', 'tie-h-k'),
    (38, 'K', 'Labor by category at contract rates, with hours', 'k-labor'),
    (39, 'K', 'Claimed indirect rate ties to Schedule A', None),
    (40, 'K', 'Direct material and ODC tie to Schedule H', None),
    (41, 'L', 'Direct labor totals tie to Schedule H', 'tie-l-h'),
    (42, 'L', 'G&A labor totals tie to Schedule B', None),
    (43, 'L', 'Other indirect pool labor totals tie to their pool schedules', None),
    (44, 'M', 'A negative statement if there is nothing to report', 'm-negative'),
    (45, 'N', 'Certificate signed at vice president or CFO level or higher', 'n-signed'),
    (46, 'O', 'Contracts listed here are shown as physically complete on Schedule I', 'tie-o-i'),
    (47, 'O', 'Level of effort, fee, period of performance and ceiling for each contract', 'o-detail'),
]

NOT_CHECKED = [
    'Whether any claimed cost is allowable, allocable or reasonable under FAR Part 31. That is judgment, not math, '
    'and DCAA audits it after adequacy.',
    'Whether the numbers match your books, payroll records and billing system. Only the workbook is checked.',
    'Checklist items marked Review, which need documents or judgment outside the workbook.',
    'Whether the certificate is actually signed: a signature normally lives on the signed copy you submit.',
    'Classified contracts, compensation cap blending agreements and anything else DCAA coordinates separately.',
]

SCOPE = ('An adequacy pre-check: are the required schedules there, does the math foot, do the schedules tie, is the '
         'certificate there, and is the submission on time. Adequacy is not allowability, and a PASS is not DCAA '
         'acceptance. DCAA makes its own adequacy determination, and the costs claimed remain your responsibility.')


# The free demo runs the three published sample workbooks (public/samples/ice-*.xlsx) and nothing else: checking your
# own submission is the paid product. Matched byte for byte; test_ice.py recomputes these from the files.
DEMO_SAMPLES = {
    '93ecdc8d21583cb74ea8f252c25b1aca949710c2728392a051df27b3845a9931': 'ice-template-clean',
    'c4ffb2020871ba4971f945cbaac0195e27d03ef554d28233e5464d9324ecd097': 'ice-own-format-clean',
    '1962cb5bbe62b6e349bfd3721aeb79d38218a7f85b88b8c648555386df597adf': 'ice-errors',
}
DEMO_ONLY = ('The free demo runs the three sample workbooks only. To check your own submission, sign in with $250.00 '
             'of credit or call the paid API. Nothing was charged.')


class InputError(ValueError):
    """The workbook cannot be checked (400, never charged). Messages never repeat workbook content."""


class Timeout(InputError):
    pass


def col_letter(i):
    s = ''
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def text(v):
    return v.strip() if isinstance(v, str) else ''


def norm(s):
    return re.sub(r'[^A-Z0-9]+', '', str(s).upper())


def rnd(x):
    return round(float(x), 2)


# ------------------------------------------------------------------ input
def parse_params(query, today=None):
    q = {str(k): (v if isinstance(v, str) else str(v)) for k, v in (query or {}).items()}
    known = {'fiscalYearEnd', 'asOf'}
    unknown = [k for k in q if k not in known]
    if unknown:
        raise InputError('Unknown parameter. Parameters are fiscalYearEnd (required) and asOf (optional), as '
                         'YYYY-MM-DD.')
    if not q.get('fiscalYearEnd'):
        raise InputError('fiscalYearEnd is required: the last day of the fiscal year this submission covers, as '
                         'YYYY-MM-DD.')

    def day(name):
        v = q.get(name, '').strip()
        try:
            return dt.date.fromisoformat(v)
        except ValueError:
            raise InputError(f'{name} must be a date written YYYY-MM-DD.') from None

    fye = day('fiscalYearEnd')
    as_of = day('asOf') if q.get('asOf') else (today or dt.datetime.now(dt.timezone.utc).date())
    return fye, as_of


def add_months_end(d, n):
    """d plus n months. A month-end date stays at month end (December 31 plus 6 months is June 30)."""
    y, m = divmod(d.month - 1 + n, 12)
    y, m = d.year + y, m + 1
    last = (dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).day
    month_end = (d + dt.timedelta(days=1)).day == 1
    return dt.date(y, m, last if month_end else min(d.day, last))


def load(body, clock):
    if not body:
        raise InputError('Upload the workbook as the request body.')
    if len(body) > MAX_BYTES:
        raise InputError(f'The workbook is {len(body) / 1e6:.1f} MB. The limit is 4 MB. Remove data tabs the '
                         'schedules do not need, or save the schedules alone as a new .xlsx, and try again.')
    if body[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        raise InputError('This is an .xls file (Excel 97 to 2003 format). This version checks .xlsx only: open it '
                         'in Excel, choose Save As, pick Excel Workbook (.xlsx), and upload that.')
    if body[:5] == b'%PDF-':
        raise InputError('This is a PDF. Upload the incurred cost submission as an .xlsx workbook.')
    if body[:2] != b'PK':
        raise InputError('This is not an .xlsx workbook. Upload the submission as an .xlsx file.')
    try:
        zf = zipfile.ZipFile(io.BytesIO(body))
        names = set(zf.namelist())
        total = sum(i.file_size for i in zf.infolist())
    except (zipfile.BadZipFile, ValueError):
        raise InputError('The file starts like an .xlsx workbook but cannot be opened.') from None
    if 'xl/workbook.xml' not in names:
        raise InputError('This ZIP file is not an .xlsx workbook.')
    if total > 200 * 1024 * 1024:
        raise InputError('The workbook expands to more than 200 MB. Upload the schedules without the data tabs.')
    import openpyxl  # noqa: PLC0415 (loaded only when a workbook arrives)
    import warnings
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any parse failure is the workbook's, and is never charged
        raise InputError('The workbook could not be read. Open it in Excel, save it as .xlsx, and try again.') from None
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True, max_col=MAX_COLS):
            rows.append(list(r))
            if len(rows) > MAX_ROWS:
                break
            if len(rows) % 500 == 0:
                clock.check()
        while rows and not any(v not in (None, '') for v in rows[-1]):
            rows.pop()
        sheets.append({'name': ws.title, 'state': getattr(ws, 'sheet_state', 'visible'), 'rows': rows})
        clock.check()
    wb.close()
    return sheets


class Clock:
    def __init__(self, budget=TIME_BUDGET_S):
        self.start = time.monotonic()
        self.budget = budget

    def check(self):
        if time.monotonic() - self.start > self.budget:
            raise Timeout('The workbook is too large to check within the time limit. Upload the schedules without '
                          'the supporting data tabs, or split very long schedules. Nothing was charged.')


def uncached_formulas(body, tabs, clock):
    """Cells in the schedule tabs that hold a formula with no saved result (a workbook written by a program and never
    opened and saved in Excel). Their values cannot be read, so the math cannot be checked."""
    import openpyxl  # noqa: PLC0415
    import warnings
    out = defaultdict(list)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    by_name = {s['name']: s for s in tabs}
    for ws in wb.worksheets:
        s = by_name.get(ws.title)
        if not s:
            continue
        for ri, r in enumerate(ws.iter_rows(values_only=True, max_col=MAX_COLS)):
            if ri >= len(s['rows']):
                break
            for ci, v in enumerate(r):
                if isinstance(v, str) and v.startswith('=') and ci < len(s['rows'][ri]) and s['rows'][ri][ci] is None:
                    out[ws.title].append(f'{col_letter(ci)}{ri + 1}')
            if ri % 500 == 0:
                clock.check()
    wb.close()
    return out


# ------------------------------------------------------------------ schedule mapping
TAB_RE = re.compile(r'^\s*(?:SCH(?:EDULE)?\.?\s*)?([A-O])(?:\s*$|\s*[-_:.)(]\s*|\s+[A-Z&(])', re.I)
SUB_RE = re.compile(r'^\s*(?:SCH(?:EDULE)?\.?\s*)?([A-O])\s*[-_ ]\s*(\d{1,2})\b|^\s*(?:SCH(?:EDULE)?\.?\s*)?([A-O])\s*\(\s*SUMMARY', re.I)
TITLE_RE = re.compile(r'\b(?:SUMMARY\s+)?SCHEDULE\s+([A-O])(?:\s*-\s*(\d{1,2}))?\b', re.I)


def title_text(rows, n=10):
    out = []
    for r in rows[:n]:
        for v in r:
            if isinstance(v, str) and v.strip():
                out.append(v.strip())
    return ' | '.join(out)


def map_schedules(sheets):
    """Each letter A to O -> its main sheet, how it was found, and any sub-schedules (F-1, H-1, H (Summary))."""
    main, subs, how = {}, defaultdict(list), {}
    claims = defaultdict(list)
    for s in sheets:
        name = s['name']
        head = title_text(s['rows'])
        sm = SUB_RE.match(name)
        if sm:
            subs[(sm.group(1) or sm.group(3)).upper()].append(s)
            continue
        m = TAB_RE.match(name)
        if m and (len(name.strip()) <= 2 or re.match(r'^\s*SCH', name, re.I) or re.match(r'^\s*[A-O]\s*[-_:.]', name)):
            claims[m.group(1).upper()].append((0, s, 'tab name'))
            continue
        t = TITLE_RE.search(head)
        if t:
            if t.group(2):
                subs[t.group(1).upper()].append(s)
            elif re.search(r'SUMMARY\s+SCHEDULE', head, re.I):
                subs[t.group(1).upper()].append(s)
            else:
                claims[t.group(1).upper()].append((1, s, 'sheet title'))
    for letter, found in claims.items():
        found.sort(key=lambda x: x[0])
        main[letter], how[letter] = found[0][1], found[0][2]
    # Content match only for letters not already found, and only on visible sheets not claimed by anything else.
    used = {id(s) for s in main.values()} | {id(x) for v in subs.values() for x in v}
    for letter in LETTERS:
        if letter in main:
            continue
        pat = re.compile(SCHEDULES[letter][1], re.I)
        cands = [s for s in sheets if id(s) not in used and s['state'] == 'visible' and pat.search(title_text(s['rows'], 8))]
        if len(cands) == 1:
            main[letter], how[letter] = cands[0], 'content'
            used.add(id(cands[0]))
        elif len(cands) > 1:
            how[letter] = 'ambiguous'
    return main, subs, how


# ------------------------------------------------------------------ table reading
TOTAL_RE = re.compile(r'(?<![A-Z_])(SUB)?TOTALS?(?![A-Z_])|GRAND\s+TOTAL', re.I)
SKIP_HEADER = re.compile(r'RATE|FACTOR|%|PERCENT|PCT|RATIO|YEAR|DATE|NUMBER|\bNO\b|CAGE|UEI|ZIP|PHONE|ACCOUNT|ORDER|KEY|'
                         r'RANK|SORT|SEQ|INDEX|LINE.?(NO|NUM)', re.I)


def is_total_label(v):
    return isinstance(v, str) and bool(TOTAL_RE.search(v.upper().replace('TOTAL_', 'TOTALX_')))


def row_kind(r):
    nums = [i for i, v in enumerate(r) if is_num(v)]
    texts = [i for i, v in enumerate(r) if isinstance(v, str) and v.strip()]
    labels = [i for i in texts if is_total_label(r[i])]
    return nums, texts, labels


class Table:
    """A sheet read as rows. Finds header rows, total rows and the header text above each column."""

    def __init__(self, sheet):
        self.name = sheet['name']
        self.rows = sheet['rows']
        self.kind = [row_kind(r) for r in self.rows]
        self.header_of = {}
        self.totals = []
        current = {}
        under_label = False      # inside the run of rows below a row of TOTAL labels
        for i, (nums, texts, labels) in enumerate(self.kind):
            r = self.rows[i]
            if not nums and not texts:
                under_label = False
                continue
            if not nums and labels and len(labels) == len(texts):
                under_label = True       # a row of TOTAL labels; the totals are on the rows below it
                continue
            if not nums and len(texts) >= 2:
                current = {c: str(r[c]).strip() for c in texts}
                under_label = False
                continue
            if nums:
                self.header_of[i] = dict(current)
                if labels or under_label:
                    self.totals.append(i)

    def header(self, i, c):
        return self.header_of.get(i, {}).get(c, '')

    def is_total(self, i):
        return i in self.totals_set

    @property
    def totals_set(self):
        if not hasattr(self, '_ts'):
            self._ts = set(self.totals)
        return self._ts

    def num(self, i, c):
        r = self.rows[i]
        return r[c] if c < len(r) and is_num(r[c]) else None

    def find_col(self, i, pattern):
        """Column whose header (for row i) matches pattern, preferring the rightmost."""
        hits = [c for c, h in self.header_of.get(i, {}).items() if re.search(pattern, h, re.I)]
        return max(hits) if hits else None


def cell(tab, i, c):
    return f'{col_letter(c)}{i + 1}'


# ------------------------------------------------------------------ report
class Report:
    def __init__(self):
        self.findings = []
        self.checks = {}

    def add(self, sev, rule, tab, rng, message, source, *, expected=None, actual=None, item=None):
        f = {'severity': sev, 'ruleId': rule, 'tab': tab, 'range': rng, 'message': message, 'source': source}
        if expected is not None:
            f['expected'] = expected
        if actual is not None:
            f['actual'] = actual
        if item:
            f['checklistItem'] = item
        self.findings.append(f)
        return f

    def mark(self, check, status):
        """status: pass, fail, warn, na, review. The worst one wins."""
        order = {'fail': 4, 'warn': 3, 'review': 2, 'pass': 1, 'na': 0}
        if check not in self.checks or order[status] > order[self.checks[check]]:
            self.checks[check] = status


# ------------------------------------------------------------------ footing
def _label_key(t, i):
    """What a total row's label says it totals, with the TOTAL words taken out ("B&P - SUBTOTAL" gives "BP")."""
    for v in t.rows[i]:
        if is_total_label(v):
            rest = re.sub(r'\b(SUB)?\s*-?\s*TOTALS?\b|\bGRAND\b|\*+|:', ' ', v.upper())
            k = norm(rest)
            if k:
                return k
    return ''


def foot_table(rep, t, letter, clock):
    """Recompute every total row from the rows above it and flag totals that do not foot.

    A total is accepted when it equals any reading of what it could total: the rows since the previous total, the rows
    sharing its key (pool, contract, category), the rows its label names, everything above it, every row in the
    table, the subtotals above it, the rows those subtotals cover, all totals above it, or the total just above it
    restated. Only when none of them agrees is it reported. One pass with running sums, so time grows with the size of
    the table, not its square."""
    if not t.totals:
        return 0
    checked = 0
    key_cols = set()
    for i in t.totals:
        nums, texts, labels = t.kind[i]
        key_cols |= {c for c in texts if c not in labels}
    key_cols = sorted(key_cols)

    def key_of(j):
        r = t.rows[j]
        return tuple(text(r[c]) if c < len(r) else '' for c in key_cols)

    total_set = t.totals_set
    detail = [i for i in range(len(t.rows)) if t.kind[i][0] and i not in total_set]
    cols = sorted({c for i in t.totals for c in t.kind[i][0]})

    def acc():
        return defaultdict(lambda: [0.0, 0])          # column -> [sum, count]

    def add(a, j):
        for c in cols:
            v = t.num(j, c)
            if v is not None:
                a[c][0] += v
                a[c][1] += 1

    grand_all = acc()
    named_idx = defaultdict(list)
    for j in detail:
        add(grand_all, j)
        for k in {norm(v) for v in t.rows[j] if isinstance(v, str) and v.strip()}:
            named_idx[k].append(j)
    named_cache = {}

    def named_sum(lk):
        if lk not in named_cache:
            a = acc()
            for j in named_idx.get(lk, ()):
                add(a, j)
            named_cache[lk] = a
        return named_cache[lk]

    before, block = acc(), acc()
    groups = defaultdict(acc)
    totals_before = acc()
    sub_stated, sub_covered = acc(), acc()       # since the last grand (non-subtotal) total
    prev_stated = None
    di = 0
    for i in t.totals:
        clock.check()
        while di < len(detail) and detail[di] < i:
            j = detail[di]
            add(before, j)
            add(block, j)
            if key_cols:
                add(groups[key_of(j)], j)
            di += 1
        grp = groups.get(key_of(i)) if key_cols else None
        lk = _label_key(t, i)
        nm = named_sum(lk) if lk else None
        cands = [block, grp, nm, before, grand_all, sub_stated, sub_covered, totals_before, prev_stated]
        for c in t.kind[i][0]:
            h = t.header(i, c)
            if not h or SKIP_HEADER.search(h):
                continue
            stated = t.rows[i][c]
            options = [a[c][0] for a in cands if a is not None and c in a and a[c][1]]
            if not options:
                continue
            checked += 1
            if any(abs(v - stated) <= TOL for v in options):
                continue
            exp_set = next(a for a in cands if a is not None and c in a and a[c][1])
            exp = exp_set[c][0]
            if exp_set is block:
                lo = max([k for k in t.totals if k < i] or [-1])
                used = [j for j in detail if lo < j < i and t.num(j, c) is not None]
            else:
                used = [j for j in detail if j < i and t.num(j, c) is not None]
            rng = f'{col_letter(c)}{used[0] + 1}:{col_letter(c)}{used[-1] + 1}' if used else cell(t, i, c)
            f = rep.add('error', 'ICE-FOOT', t.name, cell(t, i, c),
                        f'This total does not foot. The amounts above it in {rng} add up to a different figure.',
                        'dcaa-checklist', expected=rnd(exp), actual=rnd(stated))
            f['sumRange'] = rng
        # this total's own reading, for the totals that come after it
        mine = acc()
        for c in t.kind[i][0]:
            v = t.rows[i][c]
            mine[c][0] += v
            mine[c][1] += 1
            totals_before[c][0] += v
            totals_before[c][1] += 1
        if _is_sub(t, i):
            for c in t.kind[i][0]:
                sub_stated[c][0] += t.rows[i][c]
                sub_stated[c][1] += 1
            covered_by = grp if (key_cols and grp is not None and any(v[1] for v in grp.values())) else block
            for c, (sm, n) in list(covered_by.items()):
                sub_covered[c][0] += sm
                sub_covered[c][1] += n
        else:
            sub_stated, sub_covered = acc(), acc()
        prev_stated = mine
        block = acc()
    return checked


def _is_sub(t, i):
    return any(isinstance(v, str) and re.search(r'SUB[\s-]?TOTAL', v, re.I) for v in t.rows[i])


def row_math(rep, t, clock):
    """Per row: claimed = amount per books + adjustments, where a table has those three columns."""
    checked = 0
    for i, hdr in t.header_of.items():
        cols = {k: None for k in ('books', 'adj', 'claimed')}
        for c, h in hdr.items():
            H = h.upper()
            if re.search(r'ADJUST', H):
                cols['adj'] = c
            elif re.search(r'CLAIM', H) and not re.search(r'CLAIM_?TYPE', H):
                cols['claimed'] = c
            elif re.search(r'G/?L|GENERAL.?LEDGER|PER.?BOOKS?|BOOKS?|RECORDED', H) and cols['books'] is None:
                cols['books'] = c
        if None in cols.values():
            continue
        b, a, cl = (t.num(i, cols[k]) for k in ('books', 'adj', 'claimed'))
        if cl is None or (b is None and a is None):
            continue
        checked += 1
        exp = (b or 0) + (a or 0)
        if abs(exp - cl) > TOL:
            rep.add('error', 'ICE-ROW-MATH', t.name, cell(t, i, cols['claimed']),
                    f'The claimed amount should equal the amount per books plus adjustments '
                    f'({col_letter(cols["books"])}{i + 1} plus {col_letter(cols["adj"])}{i + 1}).',
                    'dcaa-checklist', expected=rnd(exp), actual=rnd(cl))
        if i % 200 == 0:
            clock.check()
    return checked


def adjustment_notes(rep, t, letter, item):
    """Rows with an adjustment need an explanatory note (checklist items 8, 12, 14, 22)."""
    missing = []
    for i, hdr in t.header_of.items():
        if i in t.totals_set:
            continue
        adj = next((c for c, h in hdr.items() if re.search(r'ADJUST', h, re.I)), None)
        note = next((c for c, h in hdr.items() if re.search(r'COMMENT|NOTE|EXPLAN|REMARK|REASON', h, re.I)), None)
        if adj is None:
            continue
        v = t.num(i, adj)
        if v is None or abs(v) < 0.5:
            continue
        r = t.rows[i]
        if note is None or note >= len(r) or not text(r[note]):
            missing.append(cell(t, i, adj))
    if missing:
        rep.add('warning', 'ICE-NOTES', t.name, ', '.join(missing[:20]) + (' and more' if len(missing) > 20 else ''),
                f'{len(missing)} {"adjustment has" if len(missing) == 1 else "adjustments have"} no explanatory note '
                'next to it. DCAA asks for a note on every adjustment and on any amount left out of the claim.',
                'dcaa-checklist', item=item)
    return bool(missing)


# ------------------------------------------------------------------ Schedule A and the ties
def pool_kind(cost_type, name):
    s = f'{cost_type} {name}'.upper()
    if re.search(r'\bCOM\b|COST OF MONEY|FCCM', s):
        return 'com'
    if re.search(r'INTER|OCCUPANCY|FACILIT|SERVICE CENTER', s):
        return 'inter'
    if re.search(r'G\s*&\s*A|\bGA\b|GENERAL', s):
        return 'ga'
    if re.search(r'FRINGE', s):
        return 'fringe'
    return 'oh'


def read_schedule_a(rep, t):
    """Rows of Schedule A: [{row, name, kind, pool, base, rate, cells}]."""
    out = []
    for i, hdr in t.header_of.items():
        if i in t.totals_set:
            continue
        pool = t.find_col(i, r'POOL.*(AMOUNT|COST|\$|EXPENSE)|^POOL$|POOL_AMOUNT|POOL AMOUNT|CLAIMED POOL')
        base = t.find_col(i, r'BASE')
        rate = t.find_col(i, r'RATE|FACTOR')
        if pool is None or base is None or rate is None:
            continue
        p, b, r = t.num(i, pool), t.num(i, base), t.num(i, rate)
        if p is None and b is None:
            continue
        name_c = next((c for c, h in hdr.items() if re.search(r'POOL.?NAME|^POOL$|DESCRIPTION|INDIRECT', h, re.I)
                       and c not in (pool, base, rate)), None)
        type_c = next((c for c, h in hdr.items() if re.search(r'TYPE', h, re.I)), None)
        row = t.rows[i]
        nm = text(row[name_c]) if name_c is not None and name_c < len(row) else ''
        ty = text(row[type_c]) if type_c is not None and type_c < len(row) else ''
        if not nm:
            nm = next((text(v) for v in row if text(v)), '')
        out.append({'row': i, 'name': nm, 'kind': pool_kind(ty, nm), 'pool': p, 'base': b, 'rate': r,
                    'cells': {'pool': cell(t, i, pool), 'base': cell(t, i, base), 'rate': cell(t, i, rate)}})
    return out


def check_rates(rep, t, rows):
    for x in rows:
        p, b, r = x['pool'], x['base'], x['rate']
        if p is None or b in (None, 0) or r is None:
            continue
        exact = p / b
        cands = [exact, exact * 100]          # a rate shown as a percentage
        dec = _decimals(r)
        tol = 0.5 * 10 ** (-dec) + 1e-9 if dec is not None else 5e-5
        if not any(abs(c - r) <= tol for c in cands):
            rep.add('error', 'ICE-RATE', t.name, x['cells']['rate'],
                    f'The rate is not the pool divided by the base ({x["cells"]["pool"]} / {x["cells"]["base"]}).',
                    'dcaa-checklist', expected=round(exact, 6), actual=r, item=1)


def _decimals(v):
    if isinstance(v, int):
        return 0
    s = repr(float(v))
    if 'e' in s:
        return None
    return min(len(s.split('.')[1]), 8) if '.' in s else 0


def pool_totals(t):
    """Claimed pool totals on a pool schedule: [(row, col, value, key text)] from its total rows, claimed column."""
    out = []
    for i in t.totals:
        nums = t.kind[i][0]
        claimed = [c for c in nums if re.search(r'CLAIM', t.header(i, c), re.I) and not re.search(r'TYPE', t.header(i, c), re.I)]
        c = max(claimed) if claimed else None
        if c is None:
            numeric_amounts = [c for c in nums if t.header(i, c) and not SKIP_HEADER.search(t.header(i, c))]
            c = max(numeric_amounts) if numeric_amounts else None
        if c is None:
            continue
        key = ' '.join(text(v) for v in t.rows[i] if isinstance(v, str))
        out.append((i, c, t.rows[i][c], key))
    return out


def best_total(totals, name):
    """The total on a pool schedule that belongs to pool `name`: matched by name, else the last (grand) total."""
    if not totals:
        return None
    n = norm(name)
    named = [x for x in totals if n and (n in norm(x[3]) or (norm(x[3]) and norm(x[3]) in n))]
    if named:
        return named[-1]
    kind = pool_kind('', name)
    by_kind = [x for x in totals if norm(x[3]) and pool_kind('', x[3]) == kind]
    if by_kind:
        return by_kind[-1]
    unkeyed = [x for x in totals if not norm(re.sub(r'(?i)\b(sub|grand)?\s*totals?\b|\bpool\b|\bexpenses?\b|\bbase\b', '', x[3]))]
    return unkeyed[-1] if unkeyed else None


def tie(rep, a_tab, a_cell, a_val, b_tab, b_cell, b_val, what, item):
    if a_val is None or b_val is None:
        return None
    ok = abs(a_val - b_val) <= TOL
    if not ok:
        rep.add('error', 'ICE-TIE', a_tab, a_cell, f'{what} does not tie: {b_tab} {b_cell} shows a different amount.',
                'dcaa-checklist', expected=rnd(b_val), actual=rnd(a_val), item=item)
        rep.add('error', 'ICE-TIE', b_tab, b_cell, f'{what} does not tie: {a_tab} {a_cell} shows a different amount.',
                'dcaa-checklist', expected=rnd(a_val), actual=rnd(b_val), item=item)
    return ok


def base_totals(t):
    """Base totals on Schedule E (or D): total rows with their key text and amount column."""
    out = []
    for i in t.totals:
        nums = [c for c in t.kind[i][0] if t.header(i, c) and not SKIP_HEADER.search(t.header(i, c))]
        if not nums:
            continue
        claimed = [c for c in nums if re.search(r'CLAIM|BASE|AMOUNT|TOTAL', t.header(i, c), re.I)]
        c = max(claimed or nums)
        key = ' '.join(text(v) for v in t.rows[i] if isinstance(v, str))
        out.append((i, c, t.rows[i][c], key))
    return out


def grand(t, pattern):
    """The value in the last total row for the column whose header matches pattern."""
    for i in reversed(t.totals):
        c = next((c for c in sorted(t.kind[i][0], reverse=True) if re.search(pattern, t.header(i, c), re.I)), None)
        if c is not None:
            return i, c, t.rows[i][c]
    return None


# ------------------------------------------------------------------ certificate, M, J, O
SIGN_FIELDS = [('signature', r'SIGNATURE'), ('name', r'NAME OF (THE )?CERTIFYING|CERTIFYING OFFICIAL'),
               ('title', r'^\s*TITLE'), ('date', r'DATE OF EXECUTION|^\s*DATE\b')]
EXEC_TITLE = re.compile(r'PRESIDENT|\bVP\b|CHIEF|\bCFO\b|\bCEO\b|\bCOO\b|OWNER|PRINCIPAL|DIRECTOR|TREASURER', re.I)


def check_certificate(rep, sheets, main):
    n = main.get('N')
    cert = None
    if n is not None:
        cert = n
    else:
        for s in sheets:
            if re.search(r'CERTIF(ICATE|Y|ICATION) .{0,40}FINAL INDIRECT|ESTABLISH FINAL INDIRECT COST RATES',
                         title_text(s['rows'], 40), re.I):
                cert = s
                break
    if cert is None:
        rep.add('error', 'ICE-CERT-MISSING', None, None, 'No Certificate of Final Indirect Costs was found (Schedule N). '
                'Without it a proposal is not accepted and the contracting officer may set the rates unilaterally.',
                'far42.703-2', item=45)
        rep.mark('certificate', 'fail')
        return False
    body = title_text(cert['rows'], 80)
    if not re.search(r'CERTIF', body, re.I):
        rep.add('error', 'ICE-CERT-MISSING', cert['name'], None, 'Schedule N was found but does not contain the '
                'certificate of final indirect costs.', 'far42.703-2', item=45)
        rep.mark('certificate', 'fail')
        return False
    blanks, title_val = [], None
    for ri, r in enumerate(cert['rows'][:80]):
        for ci, v in enumerate(r):
            if not isinstance(v, str):
                continue
            for field, pat in SIGN_FIELDS:
                if re.search(pat, v, re.I) and ':' in v:
                    after = v.split(':', 1)[1].replace('_', '').strip()
                    nxt = next((text(x) for x in r[ci + 1:] if text(x) and not re.fullmatch(r'_+', text(x))), '')
                    val = after or nxt
                    if not val:
                        blanks.append((field, f'{col_letter(ci)}{ri + 1}'))
                    elif field == 'title':
                        title_val = val
    if re.search(r'\(identify|_{8,}\s*to establish', body, re.I):
        blanks.append(('proposal', 'the first numbered paragraph'))
    if blanks:
        rep.add('warning', 'ICE-CERT-UNSIGNED', cert['name'], ', '.join(c for _, c in blanks if c[0].isalpha() and c[1:2].isdigit()) or None,
                'The certificate is there but its ' + ', '.join(sorted({f for f, _ in blanks})) + ' '
                + ('line is' if len(set(f for f, _ in blanks)) == 1 else 'lines are') + ' blank. Complete and sign it '
                'before you submit: it must be signed at vice president or CFO level or higher (FAR 52.242-4).',
                'far42.703-2', item=45)
        rep.mark('certificate', 'warn')
    elif title_val and not EXEC_TITLE.search(title_val):
        rep.add('warning', 'ICE-CERT-LEVEL', cert['name'], None, 'Check the certifying official\'s level: FAR 52.242-4 '
                'requires a vice president or chief financial officer of the business segment, or higher.',
                'far42.703-2', item=45)
        rep.mark('certificate', 'warn')
    else:
        rep.mark('certificate', 'pass')
    return True


NEGATIVE = re.compile(r'\bNONE\b|NOT APPLICABLE|\bN/?A\b|NO (DECISIONS|AGREEMENTS|CHANGES|CONTRACTS|SUBCONTRACTS|T&M|'
                      r'TIME)|NOTHING TO REPORT|THERE (WERE|ARE) NO', re.I)


def has_content(s, skip=8):
    """True when the sheet has entries below its title and header rows."""
    t = Table(s)
    header_rows = {i for i, k in enumerate(t.kind) if not k[0] and len(k[1]) >= 2}
    first_header = min(header_rows) if header_rows else skip
    for i in range(first_header + 1, len(s['rows'])):
        if i in header_rows:
            continue
        if any(v not in (None, '') for v in s['rows'][i]):
            return True
    return False


# ------------------------------------------------------------------ run
def validate(body: bytes, *, query=None, today=None, demo=False, budget=TIME_BUDGET_S):
    clock = Clock(budget)
    if demo and hashlib.sha256(body or b'').hexdigest() not in DEMO_SAMPLES:
        raise InputError(DEMO_ONLY)
    fye, as_of = parse_params(query, today)
    sheets = load(body, clock)
    rep = Report()
    main, subs, how = map_schedules(sheets)
    tables = {}

    # 1. Required schedules
    schedules = []
    for letter in LETTERS:
        s = main.get(letter)
        entry = {'schedule': letter, 'name': SCHEDULES[letter][0]}
        if s is None:
            entry.update(status='missing', tab=None, foundBy=None)
            msg = (f'Schedule {letter} ({SCHEDULES[letter][0].lower()}) was not found. ')
            if how.get(letter) == 'ambiguous':
                msg += 'More than one tab looks like it, so none was used: name the tab "' + letter + '" or "Schedule ' \
                       + letter + '".'
            else:
                msg += ('Looked for a tab named "' + letter + '" or "Schedule ' + letter + '", a sheet titled "Schedule '
                        + letter + '", or a sheet whose title describes it. If it is in the workbook under another name, '
                        'rename the tab. If it does not apply, include it with a statement that there is nothing to '
                        'report.')
            rep.add('error', 'ICE-SCHEDULE-MISSING', None, None, msg, 'far52.216-7')
            rep.mark('schedules', 'fail')
        else:
            said_none = bool(NEGATIVE.search(title_text(s['rows'], 40))) and not any(
                is_num(v) for r in s['rows'][8:] for v in r)
            empty = not has_content(s)
            st = 'found'
            if letter not in ('A', 'N'):
                st = 'not-applicable' if said_none else ('empty' if empty else 'found')
            entry.update(status=st, tab=s['name'], foundBy=how.get(letter),
                         subSchedules=[x['name'] for x in subs.get(letter, [])])
            if st == 'empty' and letter != 'M':
                rep.add('warning', 'ICE-SCHEDULE-EMPTY', s['name'], None, f'Schedule {letter} is there but empty. If '
                        'it does not apply this year, say so on the schedule (for example, "None" or "Not '
                        'applicable") so the reviewer does not read it as missing.', 'dcaa-checklist')
            tables[letter] = Table(s)
            rep.mark('schedules', 'pass')
        schedules.append(entry)
    for letter, lst in subs.items():
        for s in lst:
            tables[f'{letter}:{s["name"]}'] = Table(s)

    # 2. Formulas with no saved values
    mapped = [x for x in sheets if any(x is main.get(L) for L in LETTERS) or any(x in v for v in subs.values())]
    blank = uncached_formulas(body, mapped, clock)
    for tab, cells in blank.items():
        rep.add('error', 'ICE-NO-VALUES', tab, ', '.join(cells[:12]) + (' and more' if len(cells) > 12 else ''),
                f'{len(cells)} {"formula has" if len(cells) == 1 else "formulas have"} no saved result, so the math '
                'cannot be checked. Open the workbook in Excel, let it calculate, save it, and upload it again.',
                'dcaa-checklist')
        rep.mark('footing', 'fail')

    # 3. Footing and row math on every schedule tab
    footed = 0
    for key, t in tables.items():
        clock.check()
        footed += foot_table(rep, t, key[0], clock)
        footed += row_math(rep, t, clock)
    rep.mark('footing', 'fail' if any(f['ruleId'] in ('ICE-FOOT', 'ICE-ROW-MATH') for f in rep.findings)
             else ('pass' if footed else 'review'))
    for letter, item in (('B', 8), ('C', 12), ('D', 14), ('G', 22)):
        if letter in tables:
            adjustment_notes(rep, tables[letter], letter, item)

    # 4. Schedule A: rates, and the ties
    a_rows = []
    if 'A' in tables:
        a_rows = read_schedule_a(rep, tables['A'])
        if a_rows:
            check_rates(rep, tables['A'], a_rows)
        else:
            rep.add('error', 'ICE-A-UNREADABLE', tables['A'].name, None, 'Schedule A was found, but no rows with a '
                    'pool amount, a base and a rate could be read. Use column headings that name the pool, the base and '
                    'the rate.', 'dcaa-checklist', item=1)
    ties = run_ties(rep, tables, a_rows, clock)

    # 5. Certificate
    check_certificate(rep, sheets, main)

    # 6. Schedule-specific checklist items
    extra = checklist_specifics(rep, tables, main, a_rows)

    # 7. Deadline
    due = add_months_end(fye, 6)
    days = (due - as_of).days
    stated = stated_fye(sheets, main)
    if stated and stated != fye:
        rep.add('warning', 'ICE-FYE-MISMATCH', main['A']['name'] if 'A' in main else None, None,
                'The fiscal year end written in the workbook differs from the fiscal year end you entered. The '
                'deadline below uses the date you entered.', 'far52.216-7', expected=fye.isoformat(),
                actual=stated.isoformat())
    if days < 0:
        rep.add('warning', 'ICE-DEADLINE-PASSED', None, None,
                f'Past the 6-month mark: the submission was due by {due.isoformat()}, {-days} days before '
                f'{as_of.isoformat()}. Submit as soon as possible, unless the contracting officer granted a written '
                'extension.', 'far52.216-7', expected=due.isoformat(), actual=as_of.isoformat())
        deadline_status = 'late'
    elif days <= 30:
        deadline_status = 'due-soon'
    else:
        deadline_status = 'on-time'

    return finish(rep, schedules, ties, extra, a_rows, tables, fye, as_of, due, days, deadline_status, body, sheets)


def stated_fye(sheets, main):
    s = main.get('A') or (sheets[0] if sheets else None)
    if not s:
        return None
    m = re.search(r'FISCAL YEAR (END(ING|ED)?|END DATE)\s*[-:]?\s*(\d{1,2})/(\d{1,2})/(\d{4})', title_text(s['rows'], 10), re.I)
    if not m:
        return None
    try:
        return dt.date(int(m.group(5)), int(m.group(3)), int(m.group(4)))
    except ValueError:
        return None


def run_ties(rep, tables, a_rows, clock):
    """Checklist items 4, 5, 6, 19, 21, 41: Schedule A against B, C, D, E, F; G against H; L against H."""
    res = {}
    A = tables.get('A')

    def mark(item, ok):
        if ok is None:
            return
        res[item] = res.get(item, True) and ok

    if A and a_rows:
        b_tot = pool_totals(tables['B']) if 'B' in tables else []
        c_tot = pool_totals(tables['C']) if 'C' in tables else []
        d = tables.get('D')
        d_tot = pool_totals(d) if d else []
        e_tot = base_totals(tables['E']) if 'E' in tables else []
        for x in a_rows:
            clock.check()
            if x['kind'] == 'com':
                continue
            src = {'ga': ('B', b_tot), 'oh': ('C', c_tot), 'fringe': ('C', c_tot), 'inter': ('D', d_tot)}[x['kind']]
            tab = tables.get(src[0])
            t = best_total(src[1], x['name']) if tab and src[1] and x['pool'] is not None else None
            if t:
                mark(4, tie(rep, A.name, x['cells']['pool'], x['pool'], tab.name,
                                                   cell(tab, t[0], t[1]), t[2], 'The pool amount', 4))
            if x['base'] is None:
                continue
            if x['kind'] == 'inter' and d is not None:
                b = grand(d, r'^BASE$|BASE')
                if b:
                    mark(5, tie(rep, A.name, x['cells']['base'], x['base'], d.name, cell(d, b[0], b[1]), b[2],
                                'The intermediate pool base', 5))
            elif x['kind'] != 'inter' and e_tot and best_total(e_tot, x['name']):
                t = best_total(e_tot, x['name'])
                mark(6, tie(rep, A.name, x['cells']['base'], x['base'], tables['E'].name, cell(tables['E'], t[0], t[1]),
                            t[2], 'The allocation base', 6))
        # 19: cost of money bases on F match Schedule A
        F = tables.get('F')
        if F:
            bases = {round(x['base']) for x in a_rows if x['base'] is not None}
            def base_cells(rows):
                hits = []
                for i in rows:
                    c = next((c for c in F.kind[i][0] if re.search(r'BASE', F.header(i, c), re.I)), None)
                    if c is not None and F.rows[i][c] > 0:
                        hits.append((i, c, F.rows[i][c]))
                return hits
            col_hits = base_cells(F.totals) or base_cells([i for i in F.header_of if i not in F.totals_set])
            for i, c, v in col_hits:
                ok = any(abs(v - b) <= TOL for b in bases)
                if not ok:
                    rep.add('error', 'ICE-TIE', F.name, cell(F, i, c), 'This cost of money allocation base does not '
                            'match any base on Schedule A.', 'dcaa-checklist', actual=rnd(v), item=19)
                mark(19, ok)
    # 21: G claimed direct costs tie to H
    G, H = tables.get('G'), tables.get('H')
    hs = next((t for k, t in tables.items() if k.startswith('H:') and re.search(r'SUMMARY', t.name, re.I)), None)
    if G:
        g = grand(G, r'CLAIM')
        h = None
        if hs:
            h = grand(hs, r'GRAND.?TOTAL')
            h_tab = hs
        if h is None and H:
            h = grand(H, r'TOTAL.?DIRECT')
            h_tab = H
        if g and h:
            mark(21, tie(rep, G.name, cell(G, g[0], g[1]), g[2], h_tab.name, cell(h_tab, h[0], h[1]), h[2],
                         'Total claimed direct costs', 21))
    # 41: L direct labor ties to H labor
    L = tables.get('L')
    if L and (H or hs):
        dl = None
        for i, r in enumerate(L.rows):
            if any(isinstance(v, str) and re.fullmatch(r'\s*DIRECT LABOR\s*', v, re.I) for v in r):
                c = next((c for c in reversed(range(len(r))) if is_num(r[c])), None)
                if c is not None:
                    dl = (i, c, r[c])
                    break
        h = grand(H, r'^LABOR$|DIRECT.?LABOR') if H else None
        if dl and h:
            mark(41, tie(rep, L.name, cell(L, dl[0], dl[1]), dl[2], H.name, cell(H, h[0], h[1]), h[2],
                         'Direct labor', 41))
    return res


ID_COLS = {
    'contract': r'^(PRIME_?)?CONTRACT.?(NUMBER|NO)$|^CONTRACT$',
    'order': r'ORDER.?(NUMBER|NO)$|^ORDER$|DELIVERY.?ORDER',
    'task': r'ORDER.?TYPE|^TASK',
}
ANY_ID = r'CONTRACT.?(NUMBER|NO)|SUBCONTRACT.?(NUMBER|NO)|^CONTRACT$'


def contract_rows(t):
    """Detail rows with their contract identifiers: [(row, {'contract','order','task'}, {every id value})]."""
    out = []
    for i, hdr in t.header_of.items():
        if i in t.totals_set:
            continue
        r = t.rows[i]
        comp = {}
        for part, pat in ID_COLS.items():
            c = next((c for c, h in hdr.items() if re.search(pat, h.strip(), re.I)), None)
            if c is not None:
                comp[part] = norm(r[c]) if c < len(r) and r[c] not in (None, '') else ''
        ids = {norm(r[c]) for c, h in hdr.items() if re.search(ANY_ID, h, re.I) and c < len(r) and r[c] not in (None, '')}
        if comp.get('contract') or ids:
            out.append((i, comp, ids))
    return out


def keyed(rows, parts):
    return {j: tuple(comp.get(p, '') for p in parts) for j, comp, _ in rows}


def common_parts(a, b):
    pa = set().union(*[set(k for k, v in comp.items()) for _, comp, _ in a]) if a else set()
    pb = set().union(*[set(k for k, v in comp.items()) for _, comp, _ in b]) if b else set()
    return [p for p in ('contract', 'order', 'task') if p in pa and p in pb]


def checklist_specifics(rep, tables, main, a_rows):
    res = {}
    A = tables.get('A')
    if A:
        res[1] = bool(a_rows) and all(x['pool'] is not None and x['base'] is not None and x['rate'] is not None
                                      for x in a_rows)
    kinds = {x['kind'] for x in a_rows}
    res['present-BC'] = ('ga' not in kinds or 'B' in tables) and (not kinds & {'oh', 'fringe'} or 'C' in tables)
    res['present-D'] = 'inter' not in kinds or 'D' in tables
    D = tables.get('D')
    if D:
        hdrs = ' '.join(h for hdr in D.header_of.values() for h in hdr.values()).upper()
        res[16] = bool(re.search(r'BASE', hdrs) and re.search(r'%|PERCENT', hdrs) and re.search(r'ALLOCAT', hdrs))
        if not res[16]:
            rep.add('warning', 'ICE-D-ALLOCATION', D.name, None, 'Schedule D should show, for each recipient, the '
                    'allocation base, its percentage of the total base, and the dollars allocated.', 'dcaa-checklist',
                    item=16)
    H = tables.get('H')
    if H:
        res[23] = any(_is_sub(H, i) or i in H.totals_set for i in H.totals)
        res[26] = h_rates(rep, H, a_rows)
    I_ = tables.get('I')
    O = tables.get('O')
    K = tables.get('K')
    if I_ and H:
        res[30] = tie_i(rep, I_, H, r'GRAND.?TOTAL|^TOTAL.?COST$', r'COST|FLEX', 30)
    if I_ and K:
        res[31] = tie_i(rep, I_, K, r'TOTAL.?COST|^TOTAL$', r'T\s*&\s*M|TIME|LABOR.?HOUR', 31)
    if I_ and O:
        comp_col = None
        flagged = []
        for i, hdr in I_.header_of.items():
            c = next((c for c, h in hdr.items() if re.search(r'PHYSICALLY.?COMPLETE', h, re.I)), None)
            if c is not None and c < len(I_.rows[i]) and re.match(r'\s*(Y|YES)\b', str(I_.rows[i][c] or ''), re.I):
                flagged.append(i)
                comp_col = c
        o_ids = set().union(*[ids for _, _, ids in contract_rows(O)]) if contract_rows(O) else set()
        i_rows = {j: ids for j, _, ids in contract_rows(I_)}
        missing = [j for j in flagged if j in i_rows and not (i_rows[j] & o_ids)]
        res[33] = not missing
        if missing:
            rep.add('warning', 'ICE-I-O', I_.name, ', '.join(cell(I_, j, comp_col) for j in missing[:10]),
                    'A contract marked physically complete on Schedule I was not found on Schedule O. Check that it is '
                    'listed there, under the same contract or subcontract number.', 'dcaa-checklist', item=33)
        i_ids = set().union(*i_rows.values()) if i_rows else set()
        o_missing = [j for j, _, ids in contract_rows(O) if ids and not (ids & i_ids)]
        res[46] = not o_missing
        if o_missing:
            rep.add('warning', 'ICE-O-I', O.name, ', '.join(f'row {j + 1}' for j in o_missing[:10]),
                    'A contract on Schedule O was not found on Schedule I. Check that every contract listed as complete '
                    'is on Schedule I, shown as physically complete, under the same contract or subcontract number.',
                    'dcaa-checklist', item=46)
    if H and K:
        k_ids = set().union(*[ids for _, _, ids in contract_rows(K)]) if contract_rows(K) else set()
        missing = []
        for j, comp, ids in contract_rows(H):
            ct = next((text(v) for c, v in enumerate(H.rows[j]) if re.search(r'CONTRACT.?TYPE', H.header(j, c), re.I)), '')
            if re.search(r'T\s*&\s*M|TIME|LABOR.?HOUR', ct, re.I) and not (ids & k_ids):
                missing.append(j)
        res[37] = not missing
        if missing:
            rep.add('error', 'ICE-H-K', H.name, ', '.join(f'row {j + 1}' for j in missing[:10]),
                    'A time-and-materials or labor-hour contract on Schedule H is not on Schedule K.', 'dcaa-checklist',
                    item=37)
    if K:
        res[38] = k_labor(rep, K)
    J = tables.get('J')
    if J:
        res[35] = j_detail(rep, J)
    if O:
        hdrs = ' '.join(h for hdr in O.header_of.values() for h in hdr.values()).upper()
        need = {'level of effort': r'LOE|LEVEL OF EFFORT', 'fee': r'FEE', 'period of performance': r'PERIOD|PERFORMANCE|FROM',
                'ceiling': r'CEILING'}
        gone = [k for k, p in need.items() if not re.search(p, hdrs)]
        res[47] = not gone if has_content(main['O']) else None
        if gone and has_content(main['O']):
            rep.add('warning', 'ICE-O-DETAIL', O.name, None, 'Schedule O should show, for each contract: '
                    + ', '.join(gone) + '.', 'dcaa-checklist', item=47)
    M = main.get('M')
    if M is not None:
        res[44] = has_content(M) or bool(NEGATIVE.search(title_text(M['rows'], 60)))
        if not res[44]:
            rep.add('warning', 'ICE-M-EMPTY', M['name'], None, 'Schedule M is empty. If there are no decisions, '
                    'agreements, approvals or accounting changes to report, say so on the schedule.', 'dcaa-checklist',
                    item=44)
    return res


def h_rates(rep, H, a_rows):
    """Item 26: indirect expense on Schedule H applied at the Schedule A rates (where the columns can be identified)."""
    rate = {}
    for x in a_rows:
        if x['rate'] is None or x['pool'] is None or not x['base']:
            continue
        exact = x['pool'] / x['base']
        rate.setdefault(x['kind'] if x['kind'] != 'com' else ('com-ga' if re.search(r'G\s*&\s*A|GENERAL|\bGA\b', x['name'], re.I) else 'com-oh'), x['rate'] if abs(x['rate'] - exact) < abs(x['rate'] / 100 - exact) else x['rate'] / 100)
    pairs = [('oh', r'^OVERHEAD$|^OVERHEAD.?COST$|^OH$', r'^LABOR$|DIRECT.?LABOR'),
             ('ga', r'G&A.?COST|GA.?COST|^G&A$', r'G&A.?BASE|GA.?BASE'),
             ('com-oh', r'OVERHEAD.?COM|OH.?COM', r'^LABOR$|DIRECT.?LABOR'),
             ('com-ga', r'G&A.?COM|GA.?COM', r'G&A.?BASE|GA.?BASE')]
    checked, bad = 0, []
    for i, hdr in H.header_of.items():
        if i in H.totals_set:
            continue
        for kind, cost_p, base_p in pairs:
            if kind not in rate:
                continue
            cc = next((c for c, h in hdr.items() if re.search(cost_p, h.strip(), re.I)), None)
            bc = next((c for c, h in hdr.items() if re.search(base_p, h.strip(), re.I)), None)
            if cc is None or bc is None:
                continue
            cost, base = H.num(i, cc), H.num(i, bc)
            if cost is None or base is None:
                continue
            checked += 1
            exp = base * rate[kind]
            if abs(exp - cost) > max(TOL, abs(exp) * 0.0005):
                bad.append((i, cc, exp, cost))
    for i, c, exp, cost in bad[:30]:
        rep.add('error', 'ICE-H-RATE', H.name, cell(H, i, c), 'Indirect expense here is not the base times the claimed '
                'rate from Schedule A.', 'dcaa-checklist', expected=rnd(exp), actual=rnd(cost), item=26)
    return (not bad) if checked else None


def tie_i(rep, I_, other, total_pat, type_pat, item):
    """Items 30 and 31: current-year claimed per contract on Schedule I against Schedule H (or K), matched on the
    contract identifiers both schedules have (contract, order, task)."""
    cur = None
    for i, hdr in I_.header_of.items():
        cur = next((c for c, h in hdr.items() if re.search(r'CURRENT.?(YEAR|FY)|FY.?CLAIM|CLAIMED.?(THIS|CURRENT)', h, re.I)), None)
        if cur is not None:
            break
    if cur is None:
        return None
    i_rows, o_rows = contract_rows(I_), contract_rows(other)
    parts = common_parts(i_rows, o_rows)
    if 'contract' not in parts:
        return None
    o_key = keyed(o_rows, parts)
    o_tot = defaultdict(float)
    for j, k in o_key.items():
        c = next((c for c in sorted(other.kind[j][0], reverse=True) if re.search(total_pat, other.header(j, c), re.I)), None)
        if c is not None:
            o_tot[k] += other.rows[j][c]
    i_key = keyed(i_rows, parts)
    i_tot = defaultdict(float)
    i_first = {}
    for j, k in i_key.items():
        ct = next((text(v) for c, v in enumerate(I_.rows[j]) if re.search(r'CONTRACT.?TYPE', I_.header(j, c), re.I)), '')
        v = I_.num(j, cur)
        if not re.search(type_pat, ct, re.I) or v is None:
            continue
        i_tot[k] += v
        i_first.setdefault(k, j)
    ok_all, checked = True, 0
    for k, v in i_tot.items():
        if k not in o_tot:
            continue
        checked += 1
        if abs(v - o_tot[k]) > TOL:
            ok_all = False
            rep.add('error', 'ICE-TIE', I_.name, cell(I_, i_first[k], cur), 'This contract\'s current-year claimed costs '
                    f'do not tie to Schedule {"H" if item == 30 else "K"}.', 'dcaa-checklist', expected=rnd(o_tot[k]),
                    actual=rnd(v), item=item)
    return ok_all if checked else None


def k_labor(rep, K):
    """Item 38: labor by category with rate and hours; labor cost = rate times hours."""
    checked, bad = 0, []
    for i, hdr in K.header_of.items():
        if i in K.totals_set:
            continue
        r_c = next((c for c, h in hdr.items() if re.search(r'LABOR.?RATE|^RATE', h, re.I)), None)
        h_c = next((c for c, h in hdr.items() if re.search(r'HOURS', h, re.I)), None)
        l_c = next((c for c, h in hdr.items() if re.search(r'LABOR.?(COST|AMOUNT|DOLLARS)', h, re.I)), None)
        if None in (r_c, h_c, l_c):
            continue
        rate, hrs, cost = K.num(i, r_c), K.num(i, h_c), K.num(i, l_c)
        if None in (rate, hrs, cost):
            continue
        checked += 1
        if abs(rate * hrs - cost) > TOL:
            bad.append((i, l_c, rate * hrs, cost))
    for i, c, exp, cost in bad[:30]:
        rep.add('error', 'ICE-K-LABOR', K.name, cell(K, i, c), 'Labor cost is not the labor rate times the hours on '
                'this row.', 'dcaa-checklist', expected=rnd(exp), actual=rnd(cost), item=38)
    return (not bad) if checked else None


def j_detail(rep, J):
    """Item 35: each subcontract row has the detail DCAA lists. Only cell references are reported, never the values."""
    need = {'subcontract number': r'SUBCONTRACT.?(NUMBER|NO)', 'prime contract number': r'^CONTRACT.?(NUMBER|NO)|PRIME',
            'subcontractor name': r'NAME', 'address': r'ADDRESS', 'point of contact': r'POC|CONTACT',
            'subcontract value': r'VALUE', 'cost incurred this year': r'INCURRED|CLAIMED|COST', 'award type': r'AWARD|TYPE'}
    hdrs = {}
    for hdr in J.header_of.values():
        hdrs.update(hdr)
    if not hdrs:
        return None
    joined = ' | '.join(hdrs.values())
    gone = [k for k, p in need.items() if not re.search(p, joined, re.I)]
    if gone:
        rep.add('warning', 'ICE-J-DETAIL', J.name, None, 'Schedule J is missing columns DCAA expects for each '
                'subcontract: ' + ', '.join(gone) + '.', 'dcaa-checklist', item=35)
    return not gone


# ------------------------------------------------------------------ result
TIE_RULES = ('ICE-TIE', 'ICE-I-O', 'ICE-O-I', 'ICE-H-K')


def finish(rep, schedules, ties, extra, a_rows, tables, fye, as_of, due, days, deadline_status, body, sheets):
    order = {'error': 0, 'warning': 1}
    findings = sorted(rep.findings, key=lambda f: (order[f['severity']], f['ruleId'], f.get('tab') or '', f.get('range') or ''))
    c = Counter(f['severity'] for f in findings)
    found = {s['schedule']: s for s in schedules}
    by_item = defaultdict(list)
    for f in findings:
        if f.get('checklistItem'):
            by_item[f['checklistItem']].append(f)
    tie_items = {4: 'tie-pools', 5: 'tie-d-base', 6: 'tie-e-base', 19: 'tie-f-base', 21: 'tie-g-h', 41: 'tie-l-h'}
    checklist = []
    for item, sch, label, check in CHECKLIST:
        st = found.get(sch, {}).get('status')
        if st == 'missing':
            status = 'fail'
        elif st == 'not-applicable':
            status = 'pass' if item == 44 else 'na'
        elif check is None:
            status = 'review'
        else:
            val = ties.get(item) if item in tie_items else extra.get(item, extra.get(check))
            fs = by_item.get(item, [])
            if any(f['severity'] == 'error' for f in fs):
                status = 'fail'
            elif fs:
                status = 'warn'
            elif check.startswith('notes-'):
                status = 'pass'
            elif check == 'n-signed':
                status = {'pass': 'pass', 'warn': 'warn'}.get(rep.checks.get('certificate'), 'fail')
            elif val is True:
                status = 'pass'
            elif val is False:
                status = 'fail'
            else:
                status = 'review'
        checklist.append({'item': item, 'schedule': sch, 'label': label, 'status': status,
                          'automatic': check is not None})
    status = 'FAIL' if c['error'] else ('WARN' if c['warning'] else 'PASS')
    missing = [s['schedule'] for s in schedules if s['status'] == 'missing']
    counts = Counter(x['status'] for x in checklist)
    groups = {
        'schedules': 'fail' if missing else 'pass',
        'math': 'fail' if any(f['ruleId'] in ('ICE-FOOT', 'ICE-ROW-MATH', 'ICE-RATE', 'ICE-NO-VALUES', 'ICE-H-RATE',
                                              'ICE-K-LABOR') for f in findings) else 'pass',
        'crossTies': ('fail' if any(f['ruleId'] in TIE_RULES and f['severity'] == 'error' for f in findings) else
                      'warn' if any(f['ruleId'] in TIE_RULES for f in findings) else 'pass'),
        'certificate': rep.checks.get('certificate', 'fail'),
        'deadline': 'warn' if deadline_status == 'late' else 'pass',
    }
    return {
        'schemaVersion': 1,
        'status': status,
        'summary': ('Every required schedule is there, the math foots, the schedules tie and the certificate is in '
                    'place.' if status == 'PASS' else
                    f'{c["error"]} {"problem" if c["error"] == 1 else "problems"} to fix before you submit'
                    + (f', and {c["warning"]} to review.' if c['warning'] else '.') if c['error'] else
                    f'{c["warning"]} {"item" if c["warning"] == 1 else "items"} to review before you submit.'),
        'fiscalYearEnd': fye.isoformat(),
        'deadline': {'due': due.isoformat(), 'asOf': as_of.isoformat(), 'daysLeft': days, 'status': deadline_status,
                     'rule': 'Fiscal year end plus 6 months (FAR 52.216-7(d)(2)(i)). Extensions only in writing.'},
        'checks': groups,
        'schedules': schedules,
        'schedulesMissing': missing,
        'checklistVersion': 'DCAA Checklist for Determining Adequacy of Contractor Incurred Cost Proposal, Version 3.4, '
                            'December 2021',
        'checklist': checklist,
        'checklistCounts': {k: counts.get(k, 0) for k in ('pass', 'warn', 'fail', 'review', 'na')},
        'ratesClaimed': [{'tab': tables['A'].name, 'cells': x['cells'], 'kind': x['kind']} for x in a_rows] if 'A' in tables else [],
        'findingCount': len(findings),
        'findingCounts': {'error': c['error'], 'warning': c['warning']},
        'ruleCounts': dict(sorted(Counter(f['ruleId'] for f in findings).items())),
        'findings': findings[:MAX_FINDINGS],
        'findingsTruncated': len(findings) > MAX_FINDINGS,
        'notChecked': NOT_CHECKED,
        'sources': SOURCES,
        'scope': SCOPE,
        'input': {'format': 'xlsx', 'bytes': len(body), 'tabs': len(sheets)},
        'inputSha256': hashlib.sha256(body).hexdigest(),
    }
